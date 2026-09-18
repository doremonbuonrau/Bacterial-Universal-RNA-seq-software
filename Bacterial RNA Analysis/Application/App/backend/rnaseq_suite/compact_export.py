from __future__ import annotations

import csv
import json
import math
import os
import re
import shutil
from pathlib import Path
from typing import Any, Callable, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


EXCEL_MAX_DATA_ROWS = 1_048_575
HEADER_FILL = PatternFill("solid", fgColor="417A4B")
SUBHEADER_FILL = PatternFill("solid", fgColor="E6F2E8")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
BODY_FONT = Font(name="Calibri", size=11, color="1E2A22")
TITLE_FONT = Font(name="Calibri", size=16, bold=True, color="2A5534")
FAINT_SIDE = Side(style="thin", color="DCE5DF")
FAINT_BORDER = Border(left=FAINT_SIDE, right=FAINT_SIDE, top=FAINT_SIDE, bottom=FAINT_SIDE)


def _clean_cell(value: str, column_name: str) -> object:
    value = ILLEGAL_CHARACTERS_RE.sub("", value)
    if value == "":
        return None
    lowered = column_name.casefold()
    is_identifier = (
        lowered in {
            "gene_id", "sample_id", "run_id", "seqid", "condition", "batch",
            "role", "modality", "aligner", "status", "tool", "resolved_path",
            "version_text", "type", "path", "strand", "declared", "inferred",
        }
        or lowered.endswith("_id")
        or "sha256" in lowered
        or lowered.endswith("_md5")
        or lowered.endswith("_url")
        or lowered.endswith("_urls")
    )
    if is_identifier:
        return "'" + value if value.startswith("=") else value
    if re.fullmatch(r"[-+]?\d+", value):
        try:
            return int(value)
        except ValueError:
            pass
    if re.fullmatch(r"[-+]?(?:\d+(?:\.\d*)?|\.\d+)(?:[Ee][-+]?\d+)?", value):
        try:
            parsed = float(value)
            if math.isfinite(parsed):
                return parsed
        except ValueError:
            pass
    return "'" + value if value.startswith("=") else value


def _safe_sheet_name(value: str) -> str:
    value = re.sub(r"[\\/*?:\[\]]", "-", value).strip() or "Results"
    return value[:31]


def _unique_sheet_name(workbook: Workbook, requested: str) -> str:
    base = _safe_sheet_name(requested)
    candidate = base
    index = 2
    while candidate in workbook.sheetnames:
        suffix = f" {index}"
        candidate = base[: 31 - len(suffix)] + suffix
        index += 1
    return candidate


def _style_data_sheet(sheet, widths: list[int]) -> None:
    sheet.freeze_panes = "A2"
    if sheet.max_column and sheet.max_row:
        sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = FAINT_BORDER
    sheet.row_dimensions[1].height = 30
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = min(48, max(10, width + 2))
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = BODY_FONT
            cell.alignment = Alignment(vertical="top")
            cell.border = FAINT_BORDER
            if isinstance(cell.value, float):
                cell.number_format = "0.0000E+00" if cell.value and abs(cell.value) < 0.0001 else "0.0000"


def _write_tsv(workbook: Workbook, title: str, path: Path) -> list[str]:
    created: list[str] = []
    with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as handle:
        reader = csv.reader(handle, delimiter="\t")
        try:
            headers = next(reader)
        except StopIteration:
            return created
        headers = [str(value) for value in headers]
        sheet = workbook.create_sheet(_unique_sheet_name(workbook, title))
        created.append(sheet.title)
        sheet.append(headers)
        widths = [len(value) for value in headers]
        data_rows = 0
        chunk = 1
        for raw_row in reader:
            if data_rows >= EXCEL_MAX_DATA_ROWS:
                _style_data_sheet(sheet, widths)
                chunk += 1
                sheet = workbook.create_sheet(_unique_sheet_name(workbook, f"{title} {chunk}"))
                created.append(sheet.title)
                sheet.append(headers)
                widths = [len(value) for value in headers]
                data_rows = 0
            padded = list(raw_row) + [""] * max(0, len(headers) - len(raw_row))
            cleaned = [
                _clean_cell(padded[index], headers[index])
                for index in range(len(headers))
            ]
            sheet.append(cleaned)
            data_rows += 1
            for index, value in enumerate(cleaned):
                if value is not None:
                    widths[index] = max(widths[index], min(46, len(str(value))))
        _style_data_sheet(sheet, widths)
    return created


def _flatten(prefix: str, value: Any) -> Iterable[tuple[str, object]]:
    if isinstance(value, dict):
        for key, child in value.items():
            label = f"{prefix}.{key}" if prefix else str(key)
            yield from _flatten(label, child)
    elif isinstance(value, list):
        yield prefix, ", ".join(str(item) for item in value)
    else:
        yield prefix, value


def _write_rows(workbook: Workbook, title: str, headers: list[str], rows: list[list[object]]) -> list[str]:
    if not headers:
        return []
    sheet = workbook.create_sheet(_unique_sheet_name(workbook, title))
    sheet.append(headers)
    widths = [len(str(value)) for value in headers]
    for row in rows:
        padded = list(row) + [None] * max(0, len(headers) - len(row))
        values = padded[: len(headers)]
        sheet.append(values)
        for index, value in enumerate(values):
            if value is not None:
                widths[index] = max(widths[index], min(46, len(str(value))))
    _style_data_sheet(sheet, widths)
    return [sheet.title]


def _fastp_qc_rows(analysis_ready: Path) -> tuple[list[str], list[list[object]]]:
    headers = [
        "sample_id", "reads_before", "reads_after", "retained_percent",
        "q20_before_percent", "q30_before_percent", "q20_after_percent",
        "q30_after_percent", "adapter_trimmed_reads", "duplication_percent",
        "insert_size_peak",
    ]
    rows: list[list[object]] = []
    fastp_dir = analysis_ready / "qc" / "fastp"
    if not fastp_dir.is_dir():
        return headers, rows
    for path in sorted(fastp_dir.glob("*.fastp.json")):
        try:
            payload = json.loads(path.read_text(encoding="utf-8-sig", errors="replace"))
        except Exception:
            continue
        summary = payload.get("summary", {}) if isinstance(payload, dict) else {}
        before = summary.get("before_filtering", {}) if isinstance(summary, dict) else {}
        after = summary.get("after_filtering", {}) if isinstance(summary, dict) else {}
        adapter = payload.get("adapter_cutting", {}) if isinstance(payload, dict) else {}
        duplication = payload.get("duplication", {}) if isinstance(payload, dict) else {}
        insert_size = payload.get("insert_size", {}) if isinstance(payload, dict) else {}
        reads_before = before.get("total_reads")
        reads_after = after.get("total_reads")
        retained = None
        if isinstance(reads_before, (int, float)) and reads_before:
            retained = round(float(reads_after or 0) * 100.0 / float(reads_before), 3)
        def pct(value: object) -> object:
            if isinstance(value, (int, float)):
                return round(float(value) * 100.0, 3)
            return None
        sample_id = path.name.removesuffix(".fastp.json").replace("__run01", "")
        rows.append([
            sample_id,
            reads_before,
            reads_after,
            retained,
            pct(before.get("q20_rate")),
            pct(before.get("q30_rate")),
            pct(after.get("q20_rate")),
            pct(after.get("q30_rate")),
            adapter.get("adapter_trimmed_reads") if isinstance(adapter, dict) else None,
            pct(duplication.get("rate")) if isinstance(duplication, dict) else None,
            insert_size.get("peak") if isinstance(insert_size, dict) else None,
        ])
    return headers, rows


def _new_workbook(title: str) -> Workbook:
    workbook = Workbook()
    workbook.properties.creator = "Bacterial RNA Analysis"
    workbook.properties.title = title
    # Remove the automatically created blank sheet; the caller adds meaningful sheets.
    workbook.remove(workbook.active)
    return workbook


def _save_verified_workbook(workbook: Workbook, target: Path, required_sheets: Iterable[str]) -> Path:
    temporary = target.with_name(f".{target.stem}.tmp.xlsx")
    workbook.save(temporary)
    try:
        os.replace(temporary, target)
    except PermissionError as exc:
        temporary.unlink(missing_ok=True)
        raise RuntimeError(f"Close the existing Excel workbook before rerunning: {target}") from exc
    verification = load_workbook(target, read_only=True, data_only=False)
    try:
        missing = [name for name in required_sheets if name not in verification.sheetnames]
        if missing:
            raise RuntimeError(
                f"Excel verification failed for {target}; missing worksheet(s): {', '.join(missing)}"
            )
    finally:
        verification.close()
    return target


def _add_run_summary(workbook: Workbook, config: dict[str, Any], warnings: Iterable[str]) -> str:
    summary = workbook.create_sheet("Run summary", 0)
    summary.sheet_view.showGridLines = False
    summary.merge_cells("A1:B1")
    summary["A1"] = "RNA-seq QC and alignment results"
    summary["A1"].font = TITLE_FONT
    summary["A3"] = "Setting"
    summary["B3"] = "Value"
    for cell in summary[3]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = FAINT_BORDER
    rows = list(_flatten("configuration", config))
    rows.extend((f"warning {index}", warning) for index, warning in enumerate(warnings, start=1))
    for key, value in rows:
        summary.append([str(key), None if value is None else str(value)])
    for row in summary.iter_rows(min_row=4):
        row[0].font = Font(name="Calibri", size=11, bold=True, color="2A5534")
        row[0].fill = SUBHEADER_FILL
        row[1].font = BODY_FONT
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = FAINT_BORDER
    summary.column_dimensions["A"].width = 43
    summary.column_dimensions["B"].width = 92
    summary.freeze_panes = "A4"
    return summary.title


def _load_gene_annotation(analysis_ready: Path) -> tuple[list[str], dict[str, dict[str, str]]]:
    """Load the portable annotation table and merge gene lengths when available."""
    metadata_path = analysis_ready / "reference" / "gene_metadata.tsv"
    lengths_path = analysis_ready / "reference" / "gene_lengths.tsv"
    rows: dict[str, dict[str, str]] = {}
    headers: list[str] = []
    if metadata_path.is_file():
        with metadata_path.open("r", encoding="utf-8-sig", errors="replace", newline="") as handle:
            reader = csv.DictReader(handle, delimiter="\t")
            headers = [str(item) for item in (reader.fieldnames or [])]
            for row in reader:
                gene_id = str(row.get("gene_id", "")).strip()
                if gene_id:
                    rows[gene_id] = {str(k): str(v or "") for k, v in row.items() if k is not None}
    if lengths_path.is_file():
        with lengths_path.open("r", encoding="utf-8-sig", errors="replace", newline="") as handle:
            reader = csv.DictReader(handle, delimiter="\t")
            for row in reader:
                gene_id = str(row.get("gene_id", "")).strip()
                if not gene_id:
                    continue
                rows.setdefault(gene_id, {"gene_id": gene_id})["length_bp"] = str(row.get("length_bp", "") or "")
        if "length_bp" not in headers:
            headers.append("length_bp")
    if "gene_id" not in headers:
        headers.insert(0, "gene_id")
    return headers, rows


def _write_annotated_count_sheet(
    workbook: Workbook,
    title: str,
    count_path: Path,
    annotation_headers: list[str],
    annotation_rows: dict[str, dict[str, str]],
) -> list[str]:
    """Write raw counts with biological annotation columns directly beside gene_id."""
    with count_path.open("r", encoding="utf-8-sig", errors="replace", newline="") as handle:
        reader = csv.reader(handle, delimiter="\t")
        try:
            count_headers = [str(value) for value in next(reader)]
        except StopIteration:
            return []
        if not count_headers:
            return []
        gene_header = count_headers[0]
        sample_headers = count_headers[1:]
        annotation_order = [
            item for item in (
                "locus_tag", "gene", "name", "product", "contig", "start", "end", "strand",
                "feature_type", "original_id", "length_bp"
            ) if item in annotation_headers
        ]
        final_headers = [gene_header, *annotation_order, *sample_headers]
        sheet = workbook.create_sheet(_unique_sheet_name(workbook, title))
        sheet.append(final_headers)
        widths = [len(value) for value in final_headers]
        data_rows = 0
        chunk = 1
        for raw_row in reader:
            if data_rows >= EXCEL_MAX_DATA_ROWS:
                _style_data_sheet(sheet, widths)
                chunk += 1
                sheet = workbook.create_sheet(_unique_sheet_name(workbook, f"{title} {chunk}"))
                sheet.append(final_headers)
                widths = [len(value) for value in final_headers]
                data_rows = 0
            padded = list(raw_row) + [""] * max(0, len(count_headers) - len(raw_row))
            gene_id = str(padded[0]).strip()
            annotation = annotation_rows.get(gene_id, {})
            values: list[object] = [gene_id]
            for key in annotation_order:
                values.append(_clean_cell(str(annotation.get(key, "")), key))
            for index, header in enumerate(sample_headers, start=1):
                values.append(_clean_cell(str(padded[index]), header))
            sheet.append(values)
            data_rows += 1
            for index, value in enumerate(values):
                if value is not None:
                    widths[index] = max(widths[index], min(46, len(str(value))))
        _style_data_sheet(sheet, widths)
        sheet.freeze_panes = "A2"
        return [sheet.title]


def create_processing_workbook(
    analysis_ready: Path,
    config: dict[str, Any],
    warnings: Iterable[str],
) -> Path:
    """Create the single user-facing raw-count and annotation workbook.

    Raw integer counts remain unchanged. Gene annotation is copied beside each
    gene in the count sheet for convenient inspection, while the complete
    portable annotation and sample metadata are retained as separate sheets.
    Machine-readable TSV handoff tables remain available for automated modules.
    """

    workbook = _new_workbook("Counts & Annotation")
    created: list[str] = []
    annotation_headers, annotation_rows = _load_gene_annotation(analysis_ready)
    counts_root = analysis_ready / "counts"
    if counts_root.is_dir():
        count_paths = sorted(counts_root.rglob("*_raw_counts.tsv"))
        for path in count_paths:
            modality = path.parent.name.title()
            stem_lower = path.stem.lower()
            if stem_lower.startswith("bowtie2_"):
                sheet_title = "Bowtie2 Raw Counts"
            elif stem_lower.startswith("hisat2_"):
                sheet_title = "HISAT2 Raw Counts"
            else:
                method = path.stem.replace("_", " ").replace(" raw counts", "").strip().title()
                sheet_title = f"{modality} {method} counts"
            created.extend(
                _write_annotated_count_sheet(
                    workbook,
                    sheet_title,
                    path,
                    annotation_headers,
                    annotation_rows,
                )
            )
        # Fractional FADU output is an audit layer, not a replacement for the
        # integer featureCounts matrix, but retain it in the workbook when present.
        for path in sorted(counts_root.rglob("fadu_fractional_counts.tsv")):
            created.extend(_write_tsv(workbook, "FADU fractional audit", path))

    gene_metadata = analysis_ready / "reference" / "gene_metadata.tsv"
    gene_lengths = analysis_ready / "reference" / "gene_lengths.tsv"
    if gene_metadata.is_file():
        # Build one complete annotation worksheet with length_bp appended.
        headers = [h for h in annotation_headers if h]
        rows: list[list[object]] = []
        for gene_id in sorted(annotation_rows):
            record = annotation_rows[gene_id]
            rows.append([_clean_cell(str(record.get(h, "")), h) for h in headers])
        created.extend(_write_rows(workbook, "Gene Annotation", headers, rows))
    elif gene_lengths.is_file():
        created.extend(_write_tsv(workbook, "Gene Annotation", gene_lengths))

    sample_metadata = analysis_ready / "metadata" / "sample_metadata.tsv"
    if sample_metadata.is_file():
        created.extend(_write_tsv(workbook, "Sample Metadata", sample_metadata))

    about = workbook.create_sheet("About")
    about.sheet_view.showGridLines = False
    about["A1"] = "Bacterial RNA Analysis - Counts & Annotation"
    about["A1"].font = TITLE_FONT
    about["A3"] = "Data"
    about["B3"] = "Meaning"
    for cell in about[3]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = FAINT_BORDER
    info_rows = [
        ("Raw counts", "Unnormalized integer gene-level counts suitable for DESeq2/edgeR/limma-voom workflows."),
        ("Annotation", "Portable annotation from the matching bacterial GFF/GTF is shown beside each gene and in the Gene Annotation sheet."),
        ("Sample metadata", "Biological condition, replicate and optional batch information used by downstream modules."),
        ("Dual alignments", "When Bowtie2 + HISAT2 dual alignment is selected, the Bowtie2 and HISAT2 raw-count sheets are independent technical alternatives. Do not merge them or treat them as biological replicates."),
        ("Do not substitute", "TPM, FPKM, percentages, or transformed values are not substitutes for raw counts in count-based differential-expression significance testing."),
    ]
    for label, value in info_rows:
        about.append([label, value])
    for row in about.iter_rows(min_row=4):
        for cell in row:
            cell.border = FAINT_BORDER
            cell.font = BODY_FONT
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    about.column_dimensions["A"].width = 22
    about.column_dimensions["B"].width = 96
    created.append("About")

    count_sheet_exists = any("counts" in name.lower() for name in created)
    if not count_sheet_exists:
        raise RuntimeError("No raw-count matrix was available for the user-facing count workbook.")
    target = analysis_ready / "Counts & Annotation.xlsx"
    _save_verified_workbook(workbook, target, created)
    return target

def _inside(path: Path, root: Path) -> bool:
    try:
        path.resolve().relative_to(root.resolve())
        return True
    except ValueError:
        return False


def _delete(path: Path, root: Path, log: Callable[[str], None]) -> None:
    if not path.exists():
        return
    if not _inside(path, root) or path.resolve() == root.resolve():
        raise RuntimeError(f"Refusing cleanup outside the completed run: {path}")
    operation = "shutil.rmtree" if path.is_dir() else "Path.unlink"
    log(f"CLEANUP CODE\t{operation}({os.fspath(path)!r})")
    if path.is_dir():
        shutil.rmtree(path)
    else:
        path.unlink()


def _move_directory(
    source: Path,
    destination: Path,
    root: Path,
    log: Callable[[str], None],
) -> tuple[Path, Path] | None:
    """Move a pipeline-owned directory, safely merging an earlier partial finalization.

    A failed/resumed run may already have created the destination. In that case,
    merge the new pipeline-owned contents into it instead of aborting the run.
    """
    if not source.exists():
        return None
    if not _inside(source, root) or not _inside(destination, root):
        raise RuntimeError(f"Refusing to move a result outside the completed export: {source}")
    destination.parent.mkdir(parents=True, exist_ok=True)
    if not destination.exists():
        log(f"CLEANUP CODE\tshutil.move({os.fspath(source)!r}, {os.fspath(destination)!r})")
        shutil.move(os.fspath(source), os.fspath(destination))
        return source, destination
    if not destination.is_dir() or not source.is_dir():
        raise RuntimeError(f"Compact export path conflict: {source} -> {destination}")

    log(f"CLEANUP CODE\tmerge directory {os.fspath(source)!r} into existing {os.fspath(destination)!r}")
    for child in sorted(source.iterdir(), key=lambda item: item.name.casefold()):
        target = destination / child.name
        if child.is_dir():
            _move_directory(child, target, root, log)
            continue
        if target.exists():
            # Both locations are pipeline-owned. A resumed finalization should keep
            # the freshly generated file and replace the earlier partial copy.
            if target.is_dir():
                shutil.rmtree(target)
            else:
                target.unlink()
        log(f"CLEANUP CODE\tshutil.move({os.fspath(child)!r}, {os.fspath(target)!r})")
        shutil.move(os.fspath(child), os.fspath(target))
    try:
        source.rmdir()
    except OSError:
        pass
    return source, destination

def compact_successful_export(
    analysis_ready: Path,
    work_dir: Path,
    workbook: Path,
    log: Callable[[str], None],
) -> list[tuple[Path, Path]]:
    """Remove only verified regenerable files and group retained internals."""

    if not workbook.is_file() or workbook.stat().st_size == 0:
        raise RuntimeError("Compact cleanup was blocked because the verified Excel workbook is missing.")
    bam_root = analysis_ready / "bam"
    final_bam_root = analysis_ready / "Alignments"
    bams = []
    if bam_root.is_dir():
        bams.extend(sorted(bam_root.rglob("*.bam")))
    if final_bam_root.is_dir():
        bams.extend(sorted(final_bam_root.rglob("*.bam")))
    if not bams:
        raise RuntimeError("Compact cleanup was blocked because no final BAM files were found in bam or Alignments.")
    missing_indexes = [bam for bam in bams if not Path(str(bam) + ".bai").is_file()]
    if missing_indexes:
        raise RuntimeError(
            "Compact cleanup was blocked because BAM index files are missing for: "
            + ", ".join(path.name for path in missing_indexes)
        )

    log("")
    log("=" * 96)
    log("SAFE COMPACT OUTPUT FINALIZATION")
    log("The Excel workbook and every final BAM/BAI pair were verified before any deletion.")
    log("Only allowlisted, regenerable pipeline files are removed; unknown user files are never touched.")

    # Per-sample counter outputs are redundant after the verified matrices and
    # workbook exist. Retain only the matrices required for downstream handoff.
    counts_root = analysis_ready / "counts"
    if counts_root.is_dir():
        retained = set(counts_root.rglob("*_raw_counts.tsv"))
        retained.update(counts_root.rglob("fadu_fractional_counts.tsv"))
        for path in sorted((item for item in counts_root.rglob("*") if item.is_file()), reverse=True):
            recognized_duplicate = (
                path.name == "fadu_audit_status.tsv"
                or path.name.endswith(".featureCounts.tsv")
                or path.name.endswith(".featureCounts.tsv.summary")
                or path.name.endswith(".htseq.tsv")
                or ("fadu" in path.parts and path.name.endswith(".counts.txt"))
            )
            if path not in retained and recognized_duplicate:
                _delete(path, analysis_ready, log)
        for directory in sorted((item for item in counts_root.rglob("*") if item.is_dir()), reverse=True):
            if not any(directory.iterdir()):
                directory.rmdir()

    reference_root = analysis_ready / "reference"
    if reference_root.is_dir():
        removable_reference_names = {
            "annotation.original.gff3", "annotation.original.gtf",
            "reference.fasta.0123", "reference.fasta.amb", "reference.fasta.ann",
            "reference.fasta.bwt.2bit.64", "reference.fasta.pac", "reference.fasta.sa",
            "reference.fasta.bwt",
        }
        for path in sorted(reference_root.iterdir()):
            if path.is_file() and path.name in removable_reference_names:
                _delete(path, analysis_ready, log)

    intermediate = analysis_ready / "Intermediate files"
    intermediate.mkdir(parents=True, exist_ok=True)
    moves: list[tuple[Path, Path]] = []

    browser_files = intermediate / "Browser files"

    # Migrate folders created by an older/partially completed finalization into
    # the current compact layout. This makes finalization safe to resume after
    # a previous run stopped between directory moves.
    for legacy_source, current_destination in (
        (analysis_ready / "Reference", browser_files / "Reference"),
        (analysis_ready / "Coverage tracks", browser_files / "Coverage tracks"),
        (analysis_ready / "QC reports", intermediate / "QC reports"),
        (analysis_ready / "IGV", browser_files / "IGV"),
    ):
        moved = _move_directory(legacy_source, current_destination, analysis_ready, log)
        if moved:
            moves.append(moved)

    for source, destination in (
        (analysis_ready / "bam", analysis_ready / "Alignments"),
        (analysis_ready / "coverage", browser_files / "Coverage tracks"),
        (analysis_ready / "reference", browser_files / "Reference"),
        (analysis_ready / "counts", intermediate / "Count tables"),
        (analysis_ready / "metadata", intermediate / "Metadata"),
    ):
        moved = _move_directory(source, destination, analysis_ready, log)
        if moved:
            moves.append(moved)
        elif source == analysis_ready / "bam" and destination.is_dir():
            # A prior partial finalization may already have moved BAM/BAI files.
            # Preserve path remapping for the in-memory pipeline state.
            moves.append((source, destination))

    qc_root = analysis_ready / "qc"
    if qc_root.is_dir():
        multiqc = qc_root / "multiqc"
        moved = _move_directory(multiqc, intermediate / "QC reports", analysis_ready, log)
        if moved:
            moves.append(moved)
        summary_root = intermediate / "QC summaries"
        for source in (
            qc_root / "alignment" / "alignment_summary.tsv",
            qc_root / "strand_audit" / "strand_audit_summary.tsv",
        ):
            if source.is_file():
                destination = summary_root / source.name
                destination.parent.mkdir(parents=True, exist_ok=True)
                log(f"CLEANUP CODE\tshutil.move({os.fspath(source)!r}, {os.fspath(destination)!r})")
                shutil.move(os.fspath(source), os.fspath(destination))
                moves.append((source, destination))
        moved = _move_directory(
            qc_root,
            intermediate / "Detailed QC",
            analysis_ready,
            log,
        )
        if moved:
            moves.append(moved)

    for path in (analysis_ready / "cleaned_fastq", analysis_ready / "basecalls"):
        _delete(path, analysis_ready, log)
    if work_dir.exists():
        _delete(work_dir, work_dir.parent, log)

    log("SAFE COMPACT OUTPUT FINALIZATION COMPLETE")
    log("=" * 96)
    return moves


def remap_path(path: Path, moves: Iterable[tuple[Path, Path]]) -> Path:
    for source, destination in moves:
        try:
            relative = path.relative_to(source)
        except ValueError:
            continue
        return destination / relative
    return path
