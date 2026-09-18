#!/usr/bin/env python3
from __future__ import annotations

import argparse
import csv
import hashlib
import importlib.util
import json
import math
import os
import re
import shutil
import subprocess
import sys
import time
import urllib.parse
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, Iterator, Sequence

import requests
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from scipy.stats import hypergeom

VERSION = "1.9.77"
STRING_API = "https://version-12-0.string-db.org/api"
STRING_SPECIES_URL = "https://stringdb-downloads.org/download/species.v12.0.txt"
KEGG_API = "https://rest.kegg.jp"
DATABASE_LIBRARY_ROOT = Path(os.environ.get("BRA_DATABASE_LIBRARY", "")).expanduser() if os.environ.get("BRA_DATABASE_LIBRARY") else Path.home() / ".local" / "share" / "prok-rnaseq" / "Database Library"
_LIBRARY_ANNOUNCED = False


def ensure_database_library(*, announce: bool = True) -> Path:
    global _LIBRARY_ANNOUNCED
    DATABASE_LIBRARY_ROOT.mkdir(parents=True, exist_ok=True)
    for name in ("UniProt", "STRING", "KEGG", "Rfam", "Pathway mappings"):
        (DATABASE_LIBRARY_ROOT / name).mkdir(parents=True, exist_ok=True)
    readme = DATABASE_LIBRARY_ROOT / "README.txt"
    text = (
        "Bacterial RNA Analysis - Shared Database Library\n"
        "=================================================\n\n"
        "Downloaded and imported database resources are retained here and reused across projects.\n"
        "They are not downloaded again unless a refresh/update is explicitly requested, a cached file fails validation, or it is deleted.\n\n"
        "UniProt/ - GO and network functional annotation, Swiss-Prot/TrEMBL sequence libraries and DIAMOND indexes.\n"
        "STRING/ - cached STRING identifier/network responses keyed by organism and request.\n"
        "KEGG/ - cached organism-specific pathway mappings retrieved after explicit user confirmation.\n"
        "Rfam/ - reusable Rfam.cm/Rfam.clanin resources imported by Transcript Discovery.\n"
        "Pathway mappings/ - preserved BioCyc, MetaCyc, and custom TERM2GENE mappings.\n"
    )
    if not readme.is_file() or readme.read_text(encoding="utf-8", errors="replace") != text:
        readme.write_text(text, encoding="utf-8")
    if announce and not _LIBRARY_ANNOUNCED:
        print(f"Shared Database Library: {DATABASE_LIBRARY_ROOT}", flush=True)
        print(
            "Database resources stored here are reused automatically by later modules/projects; "
            "the same resource is not downloaded again unless it is explicitly updated or removed.",
            flush=True,
        )
        _LIBRARY_ANNOUNCED = True
    return DATABASE_LIBRARY_ROOT


def _library_manifest_entry(name: str, path: Path, metadata: dict[str, object] | None = None) -> None:
    ensure_database_library(announce=False)
    manifest_path = DATABASE_LIBRARY_ROOT / "library_manifest.json"
    try:
        manifest = json.loads(manifest_path.read_text(encoding="utf-8")) if manifest_path.is_file() else {}
    except Exception:
        manifest = {}
    entries = manifest.get("entries") if isinstance(manifest, dict) else None
    if not isinstance(entries, dict):
        entries = {}
    payload: dict[str, object] = {"path": str(path), "updated_at": time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime())}
    if path.is_file():
        payload["size_bytes"] = path.stat().st_size
    if metadata:
        payload.update(metadata)
    entries[name] = payload
    manifest = {
        "library_root": str(DATABASE_LIBRARY_ROOT), "software_version": VERSION,
        "updated_at": time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime()), "entries": entries,
    }
    temp = manifest_path.with_suffix(".json.tmp")
    temp.write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    os.replace(temp, manifest_path)


def _archive_mapping(source: Path, category: str) -> Path:
    ensure_database_library(announce=False)
    if not source.is_file():
        return source
    digest = sha256_file(source)[:16]
    target_dir = DATABASE_LIBRARY_ROOT / "Pathway mappings" / safe_name(category, "Imported")
    target_dir.mkdir(parents=True, exist_ok=True)
    suffix = source.suffix if source.suffix else ".tsv"
    target = target_dir / f"{safe_name(source.stem, 'mapping')}_{digest}{suffix}"
    if not target.is_file():
        shutil.copy2(source, target)
        print(f"Saved {category} mapping in shared Database Library for future reuse: {target}", flush=True)
    _library_manifest_entry(f"{category} mapping {digest}", target, {"sha256": sha256_file(target)})
    return target


class ExpansionError(RuntimeError):
    pass


@dataclass(frozen=True)
class Interval:
    seqid: str
    start: int  # 1-based inclusive
    end: int    # 1-based inclusive
    strand: str
    name: str = ""
    kind: str = ""
    attrs: dict[str, str] | None = None

    @property
    def length(self) -> int:
        return self.end - self.start + 1


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def safe_name(value: str, fallback: str = "item") -> str:
    cleaned = re.sub(r"[^A-Za-z0-9._-]+", "_", value.strip()).strip("._-")
    return cleaned or fallback


def parse_attrs(raw: str) -> dict[str, str]:
    raw = raw.strip()
    if not raw or raw == ".":
        return {}
    out: dict[str, str] = {}
    if "=" in raw:
        for part in raw.split(";"):
            if not part.strip():
                continue
            if "=" in part:
                k, v = part.split("=", 1)
                out[k.strip()] = urllib.parse.unquote(v.strip().strip('"'))
        return out
    for m in re.finditer(r"([^;\s]+)\s+\"([^\"]*)\"\s*;?", raw):
        out[m.group(1)] = m.group(2)
    return out


def read_gff(path: Path) -> list[Interval]:
    rows: list[Interval] = []
    with path.open("r", encoding="utf-8", errors="replace") as fh:
        for n, raw in enumerate(fh, 1):
            if not raw.strip() or raw.startswith("#"):
                continue
            fields = raw.rstrip("\r\n").split("\t")
            if len(fields) != 9:
                continue
            try:
                start, end = int(fields[3]), int(fields[4])
            except ValueError:
                continue
            attrs = parse_attrs(fields[8])
            name = (
                attrs.get("locus_tag")
                or attrs.get("gene_id")
                or attrs.get("ID")
                or attrs.get("gene")
                or attrs.get("Name")
                or f"feature_{n}"
            )
            rows.append(Interval(fields[0], start, end, fields[6], name, fields[2], attrs))
    if not rows:
        raise ExpansionError(f"No usable annotation rows were found in {path}")
    return rows


def choose_genes(rows: Sequence[Interval]) -> list[Interval]:
    for feature_type in ("gene", "CDS", "exon"):
        selected = [r for r in rows if r.kind.lower() == feature_type.lower()]
        if selected:
            return selected
    return list(rows)


def overlap_bp(a: Interval, b: Interval) -> int:
    if a.seqid != b.seqid:
        return 0
    return max(0, min(a.end, b.end) - max(a.start, b.start) + 1)


def infer_analysis_ready_paths(root: Path) -> dict[str, Path | None]:
    root = root.resolve()
    candidates: dict[str, Path | None] = {
        "fasta": None,
        "gff": None,
        "metadata": None,
        "strand_audit": None,
        "coverage_root": None,
    }
    # The 1.9.5 main pipeline publishes a minimal user-facing root and keeps
    # technical handoff data under ``intermediate``.  The older layouts remain
    # accepted so existing projects can still be opened.
    fasta_patterns = [
        "intermediate/Reference/reference.fasta",
        "intermediate/reference/reference.fasta",
        "reference/reference.fasta",
        "Reference/reference.fasta",
        "Intermediate files/Reference/reference.fasta",
    ]
    gff_patterns = [
        "intermediate/Reference/annotation.normalized.gff3",
        "intermediate/reference/annotation.normalized.gff3",
        "reference/annotation.normalized.gff3",
        "Reference/annotation.normalized.gff3",
        "Intermediate files/Reference/annotation.normalized.gff3",
    ]
    metadata_patterns = [
        "intermediate/Metadata/sample_metadata.tsv",
        "intermediate/Metadata/sample metadata.tsv",
        "intermediate/metadata/sample_metadata.tsv",
        "metadata/sample_metadata.tsv",
        "Intermediate files/Metadata/sample_metadata.tsv",
    ]
    audit_patterns = [
        "intermediate/QC/Summaries/strand_audit_summary.tsv",
        "intermediate/QC/Strand Audit/strand_audit_summary.tsv",
        "intermediate/QC/strand_audit/strand_audit_summary.tsv",
        "QC Analysis/Strand Audit/strand_audit_summary.tsv",
        "qc/strand_audit/strand_audit_summary.tsv",
        "Intermediate files/QC reports/strand_audit_summary.tsv",
    ]
    coverage_patterns = [
        "intermediate/Coverage tracks",
        "intermediate/coverage",
        "coverage",
        "Coverage tracks",
        "Browser tracks",
        "Intermediate files/Browser files",
    ]
    for key, patterns in (("fasta", fasta_patterns), ("gff", gff_patterns), ("metadata", metadata_patterns), ("strand_audit", audit_patterns)):
        for rel in patterns:
            p = root / rel
            if p.is_file():
                candidates[key] = p
                break
    for rel in coverage_patterns:
        p = root / rel
        if p.is_dir():
            candidates["coverage_root"] = p
            break
    return candidates


def parse_bedgraph(path: Path) -> Iterator[tuple[str, int, int, float]]:
    with path.open("r", encoding="utf-8", errors="replace") as fh:
        for raw in fh:
            if not raw.strip() or raw.startswith(("track", "browser", "#")):
                continue
            fields = raw.split()
            if len(fields) < 4:
                continue
            try:
                start0, end0, value = int(fields[1]), int(fields[2]), float(fields[3])
            except ValueError:
                continue
            if end0 <= start0:
                continue
            yield fields[0], start0 + 1, end0, value


def parse_bigwig(path: Path) -> Iterator[tuple[str, int, int, float]]:
    """Yield 1-based inclusive coverage intervals from a BigWig track.

    The recommended RNA Processing coverage mode intentionally keeps compact
    strand-specific BigWig rather than duplicate bedGraph files.  pyBigWig is
    therefore used only when Transcript Discovery actually needs the signal.
    """
    try:
        import pyBigWig  # type: ignore
    except ImportError as exc:
        raise ExpansionError(
            "Transcript Discovery found strand-specific BigWig tracks, but pyBigWig is missing. "
            "Run Install or repair core once to update the managed environment."
        ) from exc
    handle = pyBigWig.open(str(path))
    if handle is None:
        raise ExpansionError(f"Could not open BigWig coverage track: {path}")
    try:
        for seqid in handle.chroms():
            entries = handle.intervals(seqid) or []
            for start0, end0, value in entries:
                if end0 <= start0 or value is None:
                    continue
                yield str(seqid), int(start0) + 1, int(end0), float(value)
    finally:
        handle.close()


def parse_coverage(path: Path) -> Iterator[tuple[str, int, int, float]]:
    if path.suffix.lower() == ".bw":
        yield from parse_bigwig(path)
    else:
        yield from parse_bedgraph(path)


def segment_coverage(path: Path, strand: str, *, min_depth: float, max_gap: int, min_length: int) -> list[Interval]:
    segments: list[Interval] = []
    cur_seq = ""
    cur_start = cur_end = 0
    weighted = 0.0
    covered = 0
    idx = 0
    for seqid, start, end, value in parse_coverage(path):
        if value < min_depth:
            continue
        span = end - start + 1
        if seqid == cur_seq and cur_start and start <= cur_end + max_gap + 1:
            if start > cur_end + 1:
                gap = start - cur_end - 1
                covered += gap
            cur_end = max(cur_end, end)
            weighted += value * span
            covered += span
        else:
            if cur_start and cur_end - cur_start + 1 >= min_length:
                idx += 1
                mean_cov = weighted / max(1, covered)
                segments.append(Interval(cur_seq, cur_start, cur_end, strand, f"seg_{idx}", "transcript", {"mean_coverage": f"{mean_cov:.6g}"}))
            cur_seq, cur_start, cur_end = seqid, start, end
            weighted = value * span
            covered = span
    if cur_start and cur_end - cur_start + 1 >= min_length:
        idx += 1
        mean_cov = weighted / max(1, covered)
        segments.append(Interval(cur_seq, cur_start, cur_end, strand, f"seg_{idx}", "transcript", {"mean_coverage": f"{mean_cov:.6g}"}))
    return segments


# Backward-compatible alias used by earlier synthetic tests.
segment_bedgraph = segment_coverage


def sample_and_strand_from_path(path: Path) -> tuple[str, str] | None:
    name = path.name.lower()
    if "plus_transcript" in name:
        strand = "+"
    elif "minus_transcript" in name:
        strand = "-"
    else:
        return None
    sample = path.name.split(".raw.", 1)[0]
    if sample == path.name:
        sample = path.parent.name
    return sample, strand


def merge_sample_segments(per_sample: dict[str, list[Interval]], *, min_samples: int, min_length: int) -> list[dict[str, object]]:
    by_key: dict[tuple[str, str], list[tuple[str, Interval]]] = defaultdict(list)
    for sample, intervals in per_sample.items():
        for iv in intervals:
            by_key[(iv.seqid, iv.strand)].append((sample, iv))

    output: list[dict[str, object]] = []
    counter = 0
    for (seqid, strand), records in sorted(by_key.items()):
        records.sort(key=lambda x: (x[1].start, x[1].end))
        cluster: list[tuple[str, Interval]] = []
        cluster_end = -1
        for rec in records:
            iv = rec[1]
            if not cluster or iv.start <= cluster_end + 1:
                cluster.append(rec)
                cluster_end = max(cluster_end, iv.end)
            else:
                counter = _finalize_cluster(output, counter, seqid, strand, cluster, min_samples, min_length)
                cluster = [rec]
                cluster_end = iv.end
        if cluster:
            counter = _finalize_cluster(output, counter, seqid, strand, cluster, min_samples, min_length)
    return output


def _finalize_cluster(output: list[dict[str, object]], counter: int, seqid: str, strand: str, cluster: list[tuple[str, Interval]], min_samples: int, min_length: int) -> int:
    samples = sorted({s for s, _ in cluster})
    if len(samples) < min_samples:
        return counter
    starts = [iv.start for _, iv in cluster]
    ends = [iv.end for _, iv in cluster]
    start = int(round(sum(starts) / len(starts)))
    end = int(round(sum(ends) / len(ends)))
    if end < start:
        start, end = min(starts), max(ends)
    if end - start + 1 < min_length:
        start, end = min(starts), max(ends)
    if end - start + 1 < min_length:
        return counter
    cov_values: list[float] = []
    for _, iv in cluster:
        if iv.attrs and iv.attrs.get("mean_coverage"):
            try:
                cov_values.append(float(iv.attrs["mean_coverage"]))
            except ValueError:
                pass
    counter += 1
    output.append({
        "transcript_id": f"BRA_TX_{counter:06d}",
        "seqid": seqid,
        "start": start,
        "end": end,
        "strand": strand,
        "length": end - start + 1,
        "samples_detected": len(samples),
        "sample_ids": ",".join(samples),
        "mean_coverage": round(sum(cov_values) / len(cov_values), 6) if cov_values else "",
    })
    return counter


def classify_transcripts(transcripts: list[dict[str, object]], genes: Sequence[Interval], min_antisense_overlap: int) -> tuple[list[dict[str, object]], list[dict[str, object]]]:
    genes_by_seq: dict[str, list[Interval]] = defaultdict(list)
    for g in genes:
        genes_by_seq[g.seqid].append(g)
    for values in genes_by_seq.values():
        values.sort(key=lambda x: x.start)

    antisense: list[dict[str, object]] = []
    for tx in transcripts:
        iv = Interval(str(tx["seqid"]), int(tx["start"]), int(tx["end"]), str(tx["strand"]), str(tx["transcript_id"]))
        same_hits: list[tuple[Interval, int]] = []
        opp_hits: list[tuple[Interval, int]] = []
        for g in genes_by_seq.get(iv.seqid, []):
            if g.end < iv.start:
                continue
            if g.start > iv.end:
                break
            ov = overlap_bp(iv, g)
            if not ov:
                continue
            if g.strand == iv.strand:
                same_hits.append((g, ov))
            elif g.strand in {"+", "-"}:
                opp_hits.append((g, ov))

        tx["overlap_same_strand_genes"] = ",".join(g.name for g, _ in same_hits)
        tx["overlap_opposite_strand_genes"] = ",".join(g.name for g, _ in opp_hits)
        if same_hits:
            total_same = sum(ov for _, ov in same_hits)
            if total_same >= max(1, int(iv.length * 0.8)):
                category = "annotated_or_extension"
            else:
                category = "partial_same_strand_overlap"
        elif opp_hits:
            category = "antisense"
        else:
            category = "novel_intergenic"
        tx["category"] = category

        for gene, ov in opp_hits:
            if ov < min_antisense_overlap:
                continue
            if iv.start >= gene.start and iv.end <= gene.end:
                relationship = "enclosed"
            elif iv.strand == "+" and gene.strand == "-":
                relationship = "divergent" if iv.start <= gene.end and iv.start <= gene.start else "convergent"
            elif iv.strand == "-" and gene.strand == "+":
                relationship = "divergent" if iv.end >= gene.start and iv.end >= gene.end else "convergent"
            else:
                relationship = "partial/internal-other"
            antisense.append({
                "antisense_id": tx["transcript_id"],
                "sense_gene": gene.name,
                "seqid": iv.seqid,
                "start": iv.start,
                "end": iv.end,
                "strand": iv.strand,
                "relationship": relationship,
                "overlap_bp": ov,
                "overlap_fraction_transcript": round(ov / iv.length, 6),
                "overlap_fraction_gene": round(ov / gene.length, 6),
                "samples_detected": tx.get("samples_detected", ""),
                "mean_coverage": tx.get("mean_coverage", ""),
            })
    return transcripts, antisense


def fasta_sequences(path: Path) -> dict[str, str]:
    seqs: dict[str, list[str]] = {}
    name: str | None = None
    with path.open("r", encoding="utf-8", errors="replace") as fh:
        for raw in fh:
            line = raw.strip()
            if not line:
                continue
            if line.startswith(">"):
                name = line[1:].split()[0]
                seqs[name] = []
            elif name is not None:
                seqs[name].append(re.sub(r"\s+", "", line).upper())
    return {k: "".join(v) for k, v in seqs.items()}


def revcomp(seq: str) -> str:
    return seq.translate(str.maketrans("ACGTRYMKBDHVNacgtrymkbdhvn", "TGCAYRKMVHDBNtgcayrkmvhdbn"))[::-1]


def extract_interval_sequence(iv: Interval, seqs: dict[str, str]) -> str:
    seq = seqs.get(iv.seqid, "")
    if not seq or iv.start < 1 or iv.end > len(seq):
        return ""
    out = seq[iv.start - 1: iv.end]
    return revcomp(out) if iv.strand == "-" else out


def write_fasta(path: Path, rows: Iterable[tuple[str, str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="\n") as fh:
        for name, seq in rows:
            fh.write(f">{name}\n")
            for i in range(0, len(seq), 80):
                fh.write(seq[i:i+80] + "\n")


def parse_rnafold_output(text: str) -> dict[str, dict[str, object]]:
    results: dict[str, dict[str, object]] = {}
    current = ""
    for line in text.splitlines():
        if line.startswith(">"):
            current = line[1:].split()[0]
            results[current] = {}
        elif current and re.search(r"[().]+\s+\(\s*-?[0-9.]+\)", line):
            m = re.search(r"^([().]+)\s+\(\s*(-?[0-9.]+)\)", line.strip())
            if m:
                results[current]["dot_bracket"] = m.group(1)
                results[current]["mfe"] = float(m.group(2))
    return results


def run_command(command: Sequence[str], *, cwd: Path | None = None, stdin_path: Path | None = None, stdout_path: Path | None = None, check: bool = True) -> subprocess.CompletedProcess[str]:
    stdin = stdin_path.open("r", encoding="utf-8", errors="replace") if stdin_path else None
    stdout_handle = stdout_path.open("w", encoding="utf-8", newline="\n") if stdout_path else None
    try:
        result = subprocess.run(
            list(command),
            cwd=str(cwd) if cwd else None,
            stdin=stdin,
            stdout=stdout_handle if stdout_handle is not None else subprocess.PIPE,
            stderr=subprocess.PIPE,
            text=True,
            check=False,
        )
    finally:
        if stdin:
            stdin.close()
        if stdout_handle:
            stdout_handle.close()
    if check and result.returncode != 0:
        raise ExpansionError(f"Command failed ({result.returncode}): {' '.join(command)}\n{result.stderr[-4000:]}")
    return result


def run_rnafold(candidate_fasta: Path, output_dir: Path) -> tuple[list[dict[str, object]], str]:
    exe = shutil.which("RNAfold")
    if not exe:
        return [], "RNAfold not installed"
    result = run_command([exe, "--noPS"], stdin_path=candidate_fasta)
    raw = output_dir / "candidate_srna.rnafold.txt"
    raw.write_text(result.stdout, encoding="utf-8")
    parsed = parse_rnafold_output(result.stdout)
    rows = []
    for name, values in parsed.items():
        mfe = values.get("mfe", "")
        rows.append({"transcript_id": name, "mfe": mfe, "dot_bracket": values.get("dot_bracket", "")})
    return rows, str(raw)


def parse_rfam_tblout(path: Path) -> list[dict[str, object]]:
    rows = []
    if not path.is_file():
        return rows
    with path.open("r", encoding="utf-8", errors="replace") as fh:
        for raw in fh:
            if not raw.strip() or raw.startswith("#"):
                continue
            fields = raw.split()
            if len(fields) < 16:
                continue
            rows.append({
                "target_name": fields[0],
                "target_accession": fields[1],
                "query_name": fields[2],
                "query_accession": fields[3],
                "seq_from": fields[7],
                "seq_to": fields[8],
                "strand": fields[9],
                "score": fields[14],
                "e_value": fields[15],
            })
    return rows


def run_rfam(candidate_fasta: Path, output_dir: Path, rfam_cm: Path | None, rfam_clanin: Path | None, threads: int) -> tuple[list[dict[str, object]], str]:
    ensure_database_library()
    exe = shutil.which("cmscan")
    if not exe:
        return [], "cmscan not installed"
    rfam_dir = DATABASE_LIBRARY_ROOT / "Rfam"
    rfam_dir.mkdir(parents=True, exist_ok=True)
    library_cm = rfam_dir / "Rfam.cm"
    library_clanin = rfam_dir / "Rfam.clanin"
    if rfam_cm and rfam_cm.is_file():
        if rfam_cm.resolve() != library_cm.resolve():
            if not library_cm.is_file() or sha256_file(rfam_cm) != sha256_file(library_cm):
                shutil.copy2(rfam_cm, library_cm)
                for sibling in rfam_cm.parent.glob(rfam_cm.name + ".*"):
                    if sibling.is_file():
                        shutil.copy2(sibling, rfam_dir / sibling.name)
                print(f"Rfam covariance-model database copied into shared Database Library for future Transcript Discovery runs: {library_cm}", flush=True)
        rfam_cm = library_cm
        _library_manifest_entry("Rfam covariance models", library_cm, {"sha256": sha256_file(library_cm)})
    elif library_cm.is_file():
        rfam_cm = library_cm
        print(f"Using Rfam.cm from shared Database Library; no need to select it again: {library_cm}", flush=True)
    if rfam_clanin and rfam_clanin.is_file():
        if rfam_clanin.resolve() != library_clanin.resolve():
            if not library_clanin.is_file() or sha256_file(rfam_clanin) != sha256_file(library_clanin):
                shutil.copy2(rfam_clanin, library_clanin)
                print(f"Rfam clan file copied into shared Database Library for future Transcript Discovery runs: {library_clanin}", flush=True)
        rfam_clanin = library_clanin
        _library_manifest_entry("Rfam clan file", library_clanin, {"sha256": sha256_file(library_clanin)})
    elif library_clanin.is_file():
        rfam_clanin = library_clanin
        print(f"Using Rfam.clanin from shared Database Library; no need to select it again: {library_clanin}", flush=True)
    if not rfam_cm or not rfam_cm.is_file() or not rfam_clanin or not rfam_clanin.is_file():
        return [], f"Rfam.cm and Rfam.clanin are required the first time. After selection they are copied to {rfam_dir} and reused automatically."
    tbl = output_dir / "candidate_srna.rfam.tblout"
    command = [exe, "--cpu", str(max(1, threads)), "--cut_ga", "--rfam", "--nohmmonly", "--tblout", str(tbl), "--clanin", str(rfam_clanin), str(rfam_cm), str(candidate_fasta)]
    result = run_command(command)
    (output_dir / "candidate_srna.rfam.stdout.txt").write_text(result.stdout + "\n" + result.stderr, encoding="utf-8")
    return parse_rfam_tblout(tbl), str(tbl)


def style_workbook(wb: Workbook) -> None:
    thin = Side(style="thin", color="D9E2DD")
    header_fill = PatternFill("solid", fgColor="E7F3E9")
    header_font = Font(name="Segoe UI", size=10, bold=True, color="235B35")
    body_font = Font(name="Segoe UI", size=10)
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        if ws.max_row >= 1 and ws.max_column >= 1:
            ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.font = body_font
                cell.alignment = Alignment(vertical="top", wrap_text=False)
                cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
        for col_idx in range(1, ws.max_column + 1):
            max_len = 8
            for cell in ws.iter_cols(min_col=col_idx, max_col=col_idx, max_row=min(ws.max_row, 300)):
                for c in cell:
                    if c.value is not None:
                        max_len = max(max_len, min(60, len(str(c.value))))
            ws.column_dimensions[get_column_letter(col_idx)].width = min(62, max_len + 2)


def append_sheet(wb: Workbook, title: str, rows: list[dict[str, object]]) -> None:
    ws = wb.create_sheet(title[:31])
    if not rows:
        ws.append(["status", "message"])
        ws.append(["empty", "No rows met the selected criteria."])
        return
    fields: list[str] = []
    seen: set[str] = set()
    for row in rows:
        for key in row:
            if key not in seen:
                fields.append(key)
                seen.add(key)
    ws.append(fields)
    for row in rows:
        ws.append([row.get(k, "") for k in fields])


def write_tsv(path: Path, rows: list[dict[str, object]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if not rows:
        path.write_text("\n", encoding="utf-8")
        return
    fields = list(rows[0])
    with path.open("w", encoding="utf-8", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=fields, delimiter="\t", extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def write_transcript_gff_bed(outdir: Path, transcripts: list[dict[str, object]]) -> tuple[Path, Path]:
    gff = outdir / "novel_transcripts.gff3"
    bed = outdir / "novel_transcripts.bed"
    with gff.open("w", encoding="utf-8", newline="\n") as gh, bed.open("w", encoding="utf-8", newline="\n") as bh:
        gh.write("##gff-version 3\n")
        for tx in transcripts:
            attrs = f"ID={tx['transcript_id']};category={tx.get('category','')};sample_support={tx.get('samples_detected','')}"
            gh.write("\t".join([str(tx["seqid"]), "BacterialRNAAnalysis", "transcript", str(tx["start"]), str(tx["end"]), ".", str(tx["strand"]), ".", attrs]) + "\n")
            bh.write("\t".join([str(tx["seqid"]), str(int(tx["start"])-1), str(tx["end"]), str(tx["transcript_id"]), "0", str(tx["strand"])]) + "\n")
    return gff, bed


def _canonical_header(name: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(name).strip().lower())


def import_rockhopper_transcripts(path: Path) -> list[dict[str, object]]:
    """Import precomputed Rockhopper transcript evidence without guessing its execution CLI.

    Rockhopper distributions/exports differ in header wording. This importer intentionally
    accepts common coordinate/strand header aliases and requires the user to provide the
    already-produced table. Coordinates are treated as 1-based inclusive when numeric.
    """
    if not path.is_file():
        raise ExpansionError(f"Rockhopper transcript table not found: {path}")
    text = path.read_text(encoding="utf-8-sig", errors="replace")
    lines = [line for line in text.splitlines() if line.strip() and not line.lstrip().startswith("#")]
    if not lines:
        return []
    first = lines[0]
    delimiter = "\t" if "\t" in first else ("," if "," in first else None)
    if delimiter is None:
        # Some legacy exports are whitespace-delimited. Normalize runs of spaces to tabs.
        normalized = [re.sub(r"\s{2,}", "\t", line.strip()) for line in lines]
        delimiter = "\t"
        lines = normalized
    reader = csv.DictReader(lines, delimiter=delimiter)
    if not reader.fieldnames:
        raise ExpansionError("Rockhopper transcript table has no header row.")
    aliases = {_canonical_header(h): h for h in reader.fieldnames}
    def pick(*names: str) -> str | None:
        for name in names:
            key = _canonical_header(name)
            if key in aliases:
                return aliases[key]
        return None
    seq_col = pick("seqid", "contig", "chromosome", "chrom", "reference", "sequence")
    start_col = pick("start", "transcription start", "transcription_start", "begin", "left")
    end_col = pick("stop", "end", "transcription stop", "transcription_stop", "right")
    strand_col = pick("strand", "orientation")
    id_col = pick("transcript_id", "transcript", "name", "id")
    expr_col = pick("expression", "mean expression", "mean_expression", "rpkm", "fpkm", "tpm")
    if not seq_col or not start_col or not end_col:
        raise ExpansionError(
            "Rockhopper transcript import needs contig/sequence, start, and stop/end columns. "
            f"Detected headers: {', '.join(reader.fieldnames)}"
        )
    output: list[dict[str, object]] = []
    for idx, row in enumerate(reader, 1):
        try:
            start = int(float(str(row.get(start_col, "")).strip()))
            end = int(float(str(row.get(end_col, "")).strip()))
        except ValueError:
            continue
        if start > end:
            start, end = end, start
        raw_strand = str(row.get(strand_col, "+") if strand_col else "+").strip().lower()
        strand = "+" if raw_strand in {"+", "plus", "forward", "f", "1"} else ("-" if raw_strand in {"-", "minus", "reverse", "r", "-1"} else ".")
        output.append({
            "transcript_id": str(row.get(id_col, "") if id_col else "").strip() or f"ROCKHOPPER_TX_{idx:06d}",
            "seqid": str(row.get(seq_col, "")).strip(),
            "start": start,
            "end": end,
            "strand": strand,
            "length": end - start + 1,
            "samples_detected": "",
            "sample_ids": "",
            "mean_coverage": str(row.get(expr_col, "") if expr_col else "").strip(),
            "discovery_source": "Rockhopper imported evidence",
            "rockhopper_support": "yes",
        })
    return [r for r in output if r["seqid"] and r["strand"] in {"+", "-"}]


def merge_rockhopper_evidence(native: list[dict[str, object]], external: list[dict[str, object]]) -> list[dict[str, object]]:
    if not external:
        for row in native:
            row.setdefault("discovery_source", "native strand-coverage evidence")
            row.setdefault("rockhopper_support", "no")
        return native
    for row in native:
        row.setdefault("discovery_source", "native strand-coverage evidence")
        row.setdefault("rockhopper_support", "no")
    next_index = len(native)
    for ext in external:
        ext_iv = Interval(str(ext["seqid"]), int(ext["start"]), int(ext["end"]), str(ext["strand"]))
        best: tuple[dict[str, object], int] | None = None
        for row in native:
            if str(row.get("seqid")) != ext_iv.seqid or str(row.get("strand")) != ext_iv.strand:
                continue
            iv = Interval(ext_iv.seqid, int(row["start"]), int(row["end"]), ext_iv.strand)
            ov = overlap_bp(iv, ext_iv)
            if ov and (best is None or ov > best[1]):
                best = (row, ov)
        if best is not None:
            row, ov = best
            frac = ov / max(1, min(int(row["end"]) - int(row["start"]) + 1, ext_iv.length))
            if frac >= 0.5:
                row["rockhopper_support"] = "yes"
                row["rockhopper_transcript_id"] = ext.get("transcript_id", "")
                row["rockhopper_start"] = ext_iv.start
                row["rockhopper_end"] = ext_iv.end
                continue
        next_index += 1
        imported = dict(ext)
        imported["transcript_id"] = f"BRA_TX_{next_index:06d}"
        imported["rockhopper_original_id"] = ext.get("transcript_id", "")
        native.append(imported)
    return native


def import_generic_terminator_table(path: Path) -> list[dict[str, object]]:
    if not path.is_file():
        raise ExpansionError(f"Terminator table not found: {path}")
    text = path.read_text(encoding="utf-8-sig", errors="replace")
    lines = [line for line in text.splitlines() if line.strip() and not line.lstrip().startswith("#")]
    if not lines:
        return []
    delim = "\t" if "\t" in lines[0] else ("," if "," in lines[0] else "\t")
    reader = csv.DictReader(lines, delimiter=delim)
    return [dict(row) for row in reader if row]


def discover(args: argparse.Namespace) -> dict[str, object]:
    outdir = Path(args.output_dir).resolve()
    outdir.mkdir(parents=True, exist_ok=True)
    inferred: dict[str, Path | None] = {}
    if args.analysis_ready:
        inferred = infer_analysis_ready_paths(Path(args.analysis_ready))
    fasta = Path(args.fasta).resolve() if args.fasta else inferred.get("fasta")
    gff = Path(args.gff).resolve() if args.gff else inferred.get("gff")
    coverage_root = Path(args.coverage_root).resolve() if args.coverage_root else inferred.get("coverage_root")
    if not fasta or not Path(fasta).is_file():
        raise ExpansionError("Reference FASTA was not found. Supply --fasta or a valid analysis-ready folder.")
    if not gff or not Path(gff).is_file():
        raise ExpansionError("Gene annotation was not found. Supply --gff or a valid analysis-ready folder.")
    if not coverage_root or not Path(coverage_root).is_dir():
        raise ExpansionError("Strand-specific coverage was not found. Supply --coverage-root or use a completed RNA Processing Results folder containing intermediate/Coverage tracks.")

    # A comprehensive RNA Processing run may contain both bedGraph and BigWig
    # copies of the same signal. Select exactly one track per sample/strand,
    # preferring raw depth and then bedGraph. The recommended compact default
    # (CPM strand-specific BigWig) is supported directly.
    selected_tracks: dict[tuple[str, str], tuple[tuple[int, int], Path]] = {}
    for candidate in list(Path(coverage_root).rglob("*.bedgraph")) + list(Path(coverage_root).rglob("*.bw")):
        parsed = sample_and_strand_from_path(candidate)
        if not parsed:
            continue
        sample, strand = parsed
        lower = candidate.name.lower()
        signal_rank = 0 if ".raw." in lower else (1 if ".cpm." in lower else 2)
        format_rank = 0 if candidate.suffix.lower() == ".bedgraph" else 1
        rank = (signal_rank, format_rank)
        key = (sample, strand)
        current = selected_tracks.get(key)
        if current is None or rank < current[0]:
            selected_tracks[key] = (rank, candidate)
    coverage_files = [item[1] for item in selected_tracks.values()]
    if not coverage_files:
        raise ExpansionError(
            "No strand-specific plus_transcript/minus_transcript BigWig or bedGraph tracks were found. "
            "Transcript Discovery requires a stranded RNA Processing project or explicit strand-specific coverage."
        )

    per_sample: dict[str, list[Interval]] = defaultdict(list)
    commands = []
    for path in sorted(coverage_files):
        parsed = sample_and_strand_from_path(path)
        if not parsed:
            continue
        sample, strand = parsed
        segs = segment_coverage(path, strand, min_depth=args.min_depth, max_gap=args.max_gap, min_length=args.min_length)
        per_sample[sample].extend(segs)
        commands.append(f"native coverage segmentation {path.name}: min_depth={args.min_depth}; max_gap={args.max_gap}; min_length={args.min_length}")

    min_samples = args.min_samples
    if min_samples > len(per_sample):
        min_samples = len(per_sample)
    transcripts = merge_sample_segments(per_sample, min_samples=max(1, min_samples), min_length=args.min_length)
    rockhopper_rows: list[dict[str, object]] = []
    if args.rockhopper_transcripts:
        rockhopper_rows = import_rockhopper_transcripts(Path(args.rockhopper_transcripts).resolve())
        transcripts = merge_rockhopper_evidence(transcripts, rockhopper_rows)
        commands.append(f"import Rockhopper transcript evidence: {args.rockhopper_transcripts}")
    else:
        transcripts = merge_rockhopper_evidence(transcripts, [])
    genes = choose_genes(read_gff(Path(gff)))
    transcripts, antisense = classify_transcripts(transcripts, genes, args.min_antisense_overlap)

    seqs = fasta_sequences(Path(fasta))
    transcript_fasta_rows = []
    srna_rows: list[dict[str, object]] = []
    for tx in transcripts:
        iv = Interval(str(tx["seqid"]), int(tx["start"]), int(tx["end"]), str(tx["strand"]), str(tx["transcript_id"]))
        seq = extract_interval_sequence(iv, seqs)
        transcript_fasta_rows.append((iv.name, seq))
        if args.srna_min <= iv.length <= args.srna_max and tx.get("category") in {"novel_intergenic", "antisense"}:
            srna_rows.append({
                "transcript_id": iv.name,
                "seqid": iv.seqid,
                "start": iv.start,
                "end": iv.end,
                "strand": iv.strand,
                "length": iv.length,
                "category": tx.get("category", ""),
                "samples_detected": tx.get("samples_detected", ""),
                "mean_coverage": tx.get("mean_coverage", ""),
                "sequence": seq,
            })

    transcript_fasta = outdir / "novel_transcripts.fasta"
    srna_fasta = outdir / "candidate_srna.fasta"
    write_fasta(transcript_fasta, transcript_fasta_rows)
    write_fasta(srna_fasta, [(str(r["transcript_id"]), str(r["sequence"])) for r in srna_rows])
    gff_out, bed_out = write_transcript_gff_bed(outdir, transcripts)

    rnafold_rows: list[dict[str, object]] = []
    rfam_rows: list[dict[str, object]] = []
    notes: list[str] = []
    if args.rnafold and srna_rows:
        rnafold_rows, note = run_rnafold(srna_fasta, outdir)
        notes.append(note)
        fold_by_id = {str(r["transcript_id"]): r for r in rnafold_rows}
        for row in srna_rows:
            match = fold_by_id.get(str(row["transcript_id"]))
            if match:
                row["rnafold_mfe"] = match.get("mfe", "")
                row["rnafold_dot_bracket"] = match.get("dot_bracket", "")
    if args.rfam and srna_rows:
        rfam_rows, note = run_rfam(srna_fasta, outdir, Path(args.rfam_cm) if args.rfam_cm else None, Path(args.rfam_clanin) if args.rfam_clanin else None, args.threads)
        notes.append(note)
        hit_by_query: dict[str, list[dict[str, object]]] = defaultdict(list)
        for hit in rfam_rows:
            hit_by_query[str(hit.get("query_name", ""))].append(hit)
        for row in srna_rows:
            hits = hit_by_query.get(str(row["transcript_id"]), [])
            row["rfam_hits"] = len(hits)
            row["rfam_best_family"] = hits[0].get("target_name", "") if hits else ""
            row["rfam_best_evalue"] = hits[0].get("e_value", "") if hits else ""

    for row in srna_rows:
        score = 0
        if int(row.get("samples_detected") or 0) >= 2:
            score += 2
        if row.get("category") == "novel_intergenic":
            score += 1
        if row.get("rfam_hits", 0):
            score += 3
        if row.get("rnafold_mfe", "") != "":
            score += 1
        row["evidence_score"] = score
        row["confidence"] = "high" if score >= 5 else ("moderate" if score >= 3 else "exploratory")

    write_tsv(outdir / "predicted_transcripts.tsv", transcripts)
    write_tsv(outdir / "antisense_transcripts.tsv", antisense)
    write_tsv(outdir / "candidate_srna.tsv", srna_rows)

    wb = Workbook()
    wb.remove(wb.active)
    append_sheet(wb, "Predicted transcripts", transcripts)
    append_sheet(wb, "Rockhopper evidence", rockhopper_rows)
    append_sheet(wb, "Antisense RNAs", antisense)
    append_sheet(wb, "sRNA candidates", srna_rows)
    append_sheet(wb, "RNA structure", rnafold_rows)
    append_sheet(wb, "Rfam", rfam_rows)
    provenance = [{
        "software": "Bacterial RNA Analysis scientific expansion",
        "version": VERSION,
        "stage": "Transcript Discovery",
        "input_fasta": str(fasta),
        "input_gff": str(gff),
        "coverage_root": str(coverage_root),
        "rockhopper_transcripts": args.rockhopper_transcripts or "",
        "parameters": json.dumps({"min_depth": args.min_depth, "max_gap": args.max_gap, "min_length": args.min_length, "min_samples": args.min_samples, "min_antisense_overlap": args.min_antisense_overlap, "srna_range": [args.srna_min, args.srna_max]}, sort_keys=True),
        "notes": " | ".join(notes),
        "effective_commands": "\n".join(commands),
    }]
    append_sheet(wb, "Provenance", provenance)
    style_workbook(wb)
    workbook = outdir / "Transcript discovery and architecture.xlsx"
    wb.save(workbook)

    return {
        "status": "complete",
        "module": "transcript_discovery",
        "workbook": str(workbook),
        "transcripts": len(transcripts),
        "antisense": len(antisense),
        "srna_candidates": len(srna_rows),
        "rockhopper_evidence_rows": len(rockhopper_rows),
        "outputs": [str(workbook), str(gff_out), str(bed_out), str(transcript_fasta), str(srna_fasta)],
        "notes": notes,
    }


def read_gene_list(path: Path, column: str | None = None) -> list[str]:
    if not path.is_file():
        raise ExpansionError(f"Gene list file not found: {path}")
    aliases = {"gene", "geneid", "genename", "genenames", "genes", "id", "protein", "proteinid", "locus", "locustag"}
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            names = [name for name in wb.sheetnames if name.casefold() != "readme"]
            if not names:
                return []
            preferred = next((name for name in names if name.casefold() in {"gene list", "selected genes"}), names[0])
            rows = [list(row) for row in wb[preferred].iter_rows(values_only=True) if any(value is not None and str(value).strip() for value in row)]
        finally:
            wb.close()
    else:
        text = path.read_text(encoding="utf-8-sig", errors="replace")
        lines = [line for line in text.splitlines() if line.strip()]
        if not lines:
            return []
        delim = "\t" if "\t" in lines[0] else ("," if "," in lines[0] else None)
        rows = list(csv.reader(lines, delimiter=delim)) if delim else [[line.strip()] for line in lines]
    if not rows:
        return []
    headers = [str(value).strip() if value is not None else "" for value in rows[0]]
    keys = [_canonical_header(value) for value in headers]
    if column:
        requested = _canonical_header(column)
        if requested not in keys:
            raise ExpansionError(f"Gene column '{column}' was not found. Detected headers: {', '.join(headers)}")
        selected_index = keys.index(requested)
        start_row = 1
    else:
        selected_index = next((index for index, key in enumerate(keys) if key in aliases), 0)
        start_row = 1 if any(key in aliases for key in keys) else 0
    values = [str(row[selected_index]).strip() for row in rows[start_row:] if selected_index < len(row) and row[selected_index] is not None and str(row[selected_index]).strip()]
    return list(dict.fromkeys(values))


def request_with_retry(method: str, url: str, *, data: dict[str, str] | None = None, params: dict[str, str] | None = None, timeout: int = 60, attempts: int = 4) -> requests.Response:
    last: Exception | None = None
    for attempt in range(attempts):
        try:
            response = requests.request(method, url, data=data, params=params, timeout=timeout)
            if response.status_code in {429, 500, 502, 503, 504}:
                detail = response.text.strip().replace("\x00", "")[:1200]
                last = ExpansionError(
                    f"Remote service returned HTTP {response.status_code} for {url}"
                    + (f"\n{detail}" if detail else "")
                )
                time.sleep(min(8, 1.5 ** attempt))
                continue
            if response.status_code >= 400:
                detail = response.text.strip().replace("\x00", "")[:2000]
                raise RemoteServiceError(response.status_code, url, detail)
            return response
        except ExpansionError:
            raise
        except Exception as exc:
            last = exc
            if attempt + 1 < attempts:
                time.sleep(min(8, 1.5 ** attempt))
    raise ExpansionError(f"Request failed: {url}\n{last}")


def parse_tsv_text(text: str) -> list[dict[str, str]]:
    lines = [x for x in text.splitlines() if x.strip()]
    if not lines:
        return []
    reader = csv.DictReader(lines, delimiter="\t")
    return [dict(row) for row in reader]


def _string_species_catalog() -> list[dict[str, str]]:
    """Load the official STRING v12 organism list, retaining an offline cache."""
    ensure_database_library()
    catalog_dir = DATABASE_LIBRARY_ROOT / "STRING" / "v12"
    catalog_dir.mkdir(parents=True, exist_ok=True)
    catalog_path = catalog_dir / "species.v12.0.txt"
    if not catalog_path.is_file() or catalog_path.stat().st_size < 50:
        print("Downloading the official STRING v12 supported-organism catalog...", flush=True)
        response = request_with_retry("GET", STRING_SPECIES_URL, timeout=90)
        text = response.text
        if "taxon" not in text.casefold() or "\t" not in text:
            raise ExpansionError("The downloaded STRING organism catalog did not have the expected tabular format.")
        catalog_path.write_text(text, encoding="utf-8")
        _library_manifest_entry(
            "STRING v12 supported organisms",
            catalog_path,
            {"STRING_version": "12.0", "source": STRING_SPECIES_URL},
        )
    text = catalog_path.read_text(encoding="utf-8-sig", errors="replace")
    lines = [line for line in text.splitlines() if line.strip()]
    if not lines:
        raise ExpansionError(f"The cached STRING organism catalog is empty: {catalog_path}")
    headers = [value.lstrip("#").strip() for value in lines[0].split("\t")]
    rows: list[dict[str, str]] = []
    for values in csv.reader(lines[1:], delimiter="\t"):
        row = {headers[index]: str(value).strip() for index, value in enumerate(values) if index < len(headers)}
        canonical = {_canonical_header(key): value for key, value in row.items()}
        taxid = canonical.get("taxonid") or canonical.get("ncbitaxonid") or canonical.get("speciesid") or (values[0].strip() if values else "")
        if not taxid.isdigit():
            continue
        official = canonical.get("officialnamencbi") or canonical.get("officialname") or canonical.get("stringnamecompact") or ""
        compact = canonical.get("stringnamecompact") or official
        rows.append({"taxid": taxid, "official_name": official, "STRING_name": compact})
    if not rows:
        raise ExpansionError(f"No organisms could be parsed from the STRING catalog: {catalog_path}")
    return rows


def validate_string_taxon(taxid: int | str) -> dict[str, str]:
    value = str(taxid).strip()
    if not value.isdigit():
        raise ExpansionError("The STRING organism must be a numeric NCBI/STRING taxonomy ID.")
    match = next((row for row in _string_species_catalog() if row["taxid"] == value), None)
    if match is None:
        raise ExpansionError(
            f"Taxonomy ID {value} is not available in the official STRING v12 organism catalog. "
            "Choose a supported organism from the PPI organism list or enter another supported NCBI taxonomy ID. "
            "The software will not silently substitute a related species because that could invalidate the network."
        )
    return match


def string_species(args: argparse.Namespace) -> dict[str, object]:
    species = validate_string_taxon(args.taxid)
    print(
        f"Validated STRING organism: {species['official_name'] or species['STRING_name']} "
        f"(taxonomy ID {species['taxid']})",
        flush=True,
    )
    return {"status": "complete", "module": "string_species", **species, "STRING_version": "12.0"}


def manual_workbook(args: argparse.Namespace) -> dict[str, object]:
    helper = (
        Path(__file__).resolve().parents[2]
        / "Shared Downstream Components"
        / "Python"
        / "manual_input_workbook.py"
    )
    if not helper.is_file():
        raise ExpansionError(f"Manual Excel-input helper not found: {helper}")
    spec = importlib.util.spec_from_file_location("bra_manual_input_workbook", helper)
    if spec is None or spec.loader is None:
        raise ExpansionError(f"Could not load the manual Excel-input helper: {helper}")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    try:
        action = getattr(args, "action", "extract")
        if action == "editor-import":
            return module.import_editor(args.profile, Path(args.workbook), Path(args.output_dir))
        if action in {"editor-use", "editor-export"}:
            data = json.loads(Path(args.workbook).read_text(encoding="utf-8-sig"))
            if action == "editor-export":
                return module.export_editor(args.profile, data, Path(args.output_dir))
            return module.extract_workbook(args.profile, Path(args.workbook), Path(args.output_dir), editor_data=data)
        return module.extract_workbook(args.profile, Path(args.workbook), Path(args.output_dir))
    except Exception as exc:
        raise ExpansionError(f"Manual Excel input could not be used: {exc}") from exc


def read_expression_edges(path: Path) -> list[dict[str, object]]:
    if not path.is_file():
        raise ExpansionError(f"Expression edge table not found: {path}")
    rows: list[dict[str, object]] = []
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        values = list(ws.iter_rows(values_only=True))
        wb.close()
        if not values:
            return []
        fields = [str(value).strip() if value is not None else "" for value in values[0]]
        rows = [
            {fields[index]: value for index, value in enumerate(row) if index < len(fields) and fields[index]}
            for row in values[1:]
            if any(value is not None and str(value).strip() for value in row)
        ]
    else:
        text = path.read_text(encoding="utf-8-sig", errors="replace")
        lines = [line for line in text.splitlines() if line.strip()]
        if not lines:
            return []
        delim = "\t" if "\t" in lines[0] else ","
        reader = csv.DictReader(lines, delimiter=delim)
        fields = reader.fieldnames or []
        rows = list(reader)
    lower = {_canonical_header(f): f for f in fields}
    def field(*names: str) -> str | None:
        for name in names:
            key = _canonical_header(name)
            if key in lower:
                return lower[key]
        return None
    source_col = field("source", "from", "gene1", "node1", "regulator")
    target_col = field("target", "to", "gene2", "node2", "regulated")
    weight_col = field("weight", "correlation", "score", "importance", "edge_weight")
    if not source_col or not target_col:
        raise ExpansionError(f"Expression edge table needs source/target columns. Detected: {', '.join(fields)}")
    out: list[dict[str, object]] = []
    for row in rows:
        a, b = str(row.get(source_col) or "").strip(), str(row.get(target_col) or "").strip()
        if not a or not b:
            continue
        out.append({
            "source": a,
            "target": b,
            "edge_type": "expression_network",
            "expression_weight": str(row.get(weight_col, "")).strip() if weight_col else "",
        })
    return out


def combined_network_evidence(string_edges: list[dict[str, object]], expression_edges: list[dict[str, object]]) -> list[dict[str, object]]:
    pairs: dict[tuple[str, str], dict[str, object]] = {}
    for edge in expression_edges:
        a, b = sorted((str(edge["source"]), str(edge["target"])))
        row = pairs.setdefault((a, b), {"source": a, "target": b, "expression_support": "no", "STRING_support": "no", "supporting_layers": 0})
        row["expression_support"] = "yes"
        row["expression_weight"] = edge.get("expression_weight", "")
    for edge in string_edges:
        a, b = sorted((str(edge["source"]), str(edge["target"])))
        row = pairs.setdefault((a, b), {"source": a, "target": b, "expression_support": "no", "STRING_support": "no", "supporting_layers": 0})
        row["STRING_support"] = "yes"
        row["STRING_combined_score"] = edge.get("combined_score", "")
        row["STRING_network_type"] = edge.get("edge_type", "")
    for row in pairs.values():
        row["supporting_layers"] = int(row.get("expression_support") == "yes") + int(row.get("STRING_support") == "yes")
        weights: list[float] = []
        for key in ("STRING_combined_score", "expression_weight"):
            try:
                value = abs(float(row.get(key, "")))
            except (TypeError, ValueError):
                continue
            # STRING REST scores are normally 0-1, while imported tables can
            # use the documented 0-1000 confidence scale.
            weights.append(value / 1000.0 if value > 1.0 else value)
        row["weight"] = max(weights, default=1.0)
    return sorted(pairs.values(), key=lambda r: (-int(r["supporting_layers"]), str(r["source"]), str(r["target"])))



class RemoteServiceError(ExpansionError):
    def __init__(self, status: int, url: str, detail: str):
        self.status, self.url, self.detail = status, url, detail
        super().__init__(f"Remote service rejected the request (HTTP {status}) at {url}.\nService response: {detail}")


def string_aliases(genes: list[str], paths: list[Path]) -> dict[str, list[str]]:
    """Only use explicit annotation relationships; never infer aliases from tag spelling."""
    from urllib.parse import unquote
    wanted = set(genes)
    aliases: dict[str, list[str]] = {gene: [] for gene in genes}
    for path in dict.fromkeys(paths):
        if not path.is_file():
            continue
        records = []
        if path.suffix.lower() in {".gff", ".gff3", ".gtf"}:
            for line in path.read_text(encoding="utf-8-sig", errors="replace").splitlines():
                if line.startswith("#"):
                    continue
                fields = line.split("\t")
                if len(fields) < 9:
                    continue
                row = {}
                for part in fields[8].split(";"):
                    if "=" in part:
                        key, value = part.strip().split("=", 1)
                    elif " " in part.strip():
                        key, value = part.strip().split(" ", 1)
                    else:
                        continue
                    row[_canonical_header(key)] = unquote(value.strip(' "'))
                records.append(row)
        elif path.suffix.lower() in {".xlsx", ".xlsm"}:
            wb = load_workbook(path, read_only=True, data_only=True)
            try:
                for ws in wb:
                    if ws.title.lower().startswith(("readme", "example")):
                        continue
                    values = iter(ws.values)
                    headers = [_canonical_header(str(v or "")) for v in next(values, [])]
                    records.extend(dict(zip(headers, values)) for values in values)
            finally:
                wb.close()
        else:
            text = path.read_text(encoding="utf-8-sig", errors="replace")
            if not text.strip():
                continue
            delim = "\t" if "\t" in text.splitlines()[0] else ","
            records = [{_canonical_header(k): v for k, v in row.items() if k} for row in csv.DictReader(text.splitlines(), delimiter=delim)]
        for row in records:
            keys = [str(row.get(k) or "").strip() for k in ("geneid", "locustag", "oldlocustag", "id", "gene", "parent")]
            matches = wanted.intersection(keys)
            if not matches:
                continue
            candidates = []
            for key in ("proteinid", "proteinaccession", "uniprot", "uniprotid", "uniprotaccession", "stringid", "alias", "oldlocustag", "gene", "genename"):
                for value in re.split(r"[;,|]", str(row.get(key) or "")):
                    value = value.strip()
                    if value and not any(c.isspace() for c in value):
                        candidates.append(value)
            for item in str(row.get("dbxref") or "").split(","):
                if ":" in item:
                    db, value = item.split(":", 1)
                    if db.lower() in {"uniprotkb", "uniprotkb/swiss-prot", "uniprotkb/trembl", "refseq"}:
                        candidates.append(value.strip())
            for gene in matches:
                for candidate in candidates:
                    if candidate != gene and candidate not in aliases[gene]:
                        aliases[gene].append(candidate)
    return aliases


def string_mapping_query(identifiers: list[str], taxid: int, tax_dir: Path) -> list[dict[str, str]]:
    result = []
    for start in range(0, len(identifiers), 1000):
        batch = identifiers[start:start + 1000]
        key = hashlib.sha256(json.dumps({"ids": batch, "taxid": taxid, "echo_query": 1, "schema": 2}, sort_keys=True).encode()).hexdigest()[:20]
        cache = tax_dir / f"mapping_{key}.tsv"
        if cache.is_file():
            text = cache.read_text(encoding="utf-8")
        else:
            try:
                time.sleep(1)  # STRING requests must be spaced at least one second apart.
                response = request_with_retry("POST", f"{STRING_API}/tsv/get_string_ids", data={
                    "identifiers": "\r".join(batch), "species": str(taxid), "echo_query": "1", "caller_identity": "Bacterial_RNA_Analysis"})
                text = response.text
            except RemoteServiceError as exc:
                if exc.status == 404 and ("nothing found" in exc.detail.lower() or "did not find any matches" in exc.detail.lower()):
                    print(f"STRING found no identifier matches in this batch ({len(batch)} identifiers).", flush=True)
                    continue
                raise
            cache.write_text(text, encoding="utf-8")
        for row in parse_tsv_text(text):
            query = row.get("queryItem", "")
            if not query:
                try:
                    index = int(row.get("queryIndex", "-1"))
                    query = batch[index] if 0 <= index < len(batch) else ""
                except (TypeError, ValueError):
                    pass
            if query not in batch or not row.get("stringId"):
                continue
            if row.get("ncbiTaxonId") and str(row["ncbiTaxonId"]) != str(taxid):
                raise ExpansionError("STRING returned a mapping for a different organism; retrieval stopped.")
            result.append({**row, "queryItem": query})
    return result


def write_string_audit(outdir: Path, audit: list[dict], unmatched: list[dict]) -> None:
    folder = outdir / "Intermediate files"
    folder.mkdir(parents=True, exist_ok=True)
    for name, rows, fields in (("STRING identifier mapping.tsv", audit, ["input_id", "queryItem", "stringId", "preferredName", "mapping_source"]),
                               ("STRING unmapped identifiers.tsv", unmatched, ["input_id", "status", "aliases_tried"])):
        with (folder / name).open("w", encoding="utf-8", newline="") as handle:
            writer = csv.DictWriter(handle, delimiter="\t", fieldnames=fields, extrasaction="ignore")
            writer.writeheader(); writer.writerows(rows)


def discover_string_alias_paths(gene_list: Path, outdir: Path, explicit: Path | None = None) -> list[Path]:
    """Find matching project annotations without guessing identifiers or species.

    DE exports are commonly nested two or three folders below the RNA-processing
    GFF. Search only nearby project roots, use filename/type evidence, and leave
    the biological mapping itself to ``string_aliases`` (which accepts only
    explicit same-row GFF/table relationships).
    """
    found: list[tuple[int, Path]] = []
    seen: set[Path] = set()

    def add(path: Path, score: int) -> None:
        try:
            resolved = path.resolve()
        except OSError:
            return
        if resolved in seen or not resolved.is_file():
            return
        seen.add(resolved)
        found.append((score, resolved))

    if explicit is not None:
        add(explicit, 10_000)

    exact_names = {
        "identifier_aliases.tsv": 950,
        "gene_aliases.tsv": 930,
        "gene_annotation.tsv": 900,
        "annotation_mapping_used.tsv": 880,
        "annotation.tsv": 820,
    }
    roots: list[Path] = []
    # Limit traversal to the file's first three ancestors (typically GO input,
    # DE analysis, and project root) plus the chosen result directory. Never
    # walk an entire mounted Windows drive.
    for candidate in [*list(gene_list.parents)[:3], outdir]:
        try:
            resolved = candidate.resolve()
        except OSError:
            continue
        # A shallow input such as /tmp/genes.tsv must not turn /tmp or / into a
        # project search root. Likewise, /mnt/e is a complete Windows volume,
        # not a project. The selected output folder remains eligible.
        parts = resolved.parts
        is_volume_root = len(parts) == 3 and len(parts[-1]) == 1 and parts[-2].casefold() == "mnt"
        if len(parts) <= 2 or is_volume_root:
            continue
        if resolved.is_dir() and resolved not in roots:
            roots.append(resolved)

    ignored_directories = {".git", "__pycache__", "node_modules", "string cache"}
    for root in roots:
        root_depth = len(root.parts)
        visited = 0
        for current, directories, filenames in os.walk(root):
            depth = len(Path(current).parts) - root_depth
            directories[:] = [
                name for name in directories
                if name.casefold() not in ignored_directories and not name.startswith(".") and depth < 5
            ]
            for filename in filenames:
                visited += 1
                if visited > 30_000:
                    directories[:] = []
                    break
                lower = filename.casefold()
                path = Path(current) / filename
                if lower in {"string identifier mapping.tsv", "string unmapped identifiers.tsv"}:
                    continue
                suffix = path.suffix.casefold()
                if suffix in {".gff", ".gff3", ".gtf"}:
                    rank = 980 + (40 if re.search(r"reference|genomic|annotation", lower) else 0) - depth * 3
                    add(path, rank)
                    continue
                if suffix not in {".tsv", ".csv", ".xlsx", ".xlsm"}:
                    continue
                rank = exact_names.get(lower, 0)
                if not rank and re.search(r"identifier.*alias|gene.*(?:protein|uniprot)|annotation.*mapping|gene.*annotation", lower):
                    rank = 760
                if rank:
                    add(path, rank - depth * 3)
            if visited > 30_000:
                break
    return [path for _score, path in sorted(found, key=lambda item: (-item[0], len(item[1].parts), str(item[1]).casefold()))[:40]]


def string_network(args: argparse.Namespace) -> dict[str, object]:
    ensure_database_library()
    if not 0 <= int(args.required_score) <= 1000:
        raise ExpansionError("Minimum STRING score must be an integer from 0 through 1000.")
    if int(args.add_nodes) < 0:
        raise ExpansionError("The number of added STRING interactors cannot be negative.")
    species = validate_string_taxon(args.taxid)
    print(
        f"Validated STRING organism: {species['official_name'] or species['STRING_name']} "
        f"(taxonomy ID {species['taxid']}).",
        flush=True,
    )
    genes = read_gene_list(Path(args.gene_list), args.gene_column)
    if not genes:
        raise ExpansionError("No gene/protein identifiers were found.")
    outdir = Path(args.output_dir).resolve()
    cache = outdir / "Intermediate files" / "STRING cache"
    cache.mkdir(parents=True, exist_ok=True)
    outdir.mkdir(parents=True, exist_ok=True)

    tax_dir = DATABASE_LIBRARY_ROOT / "STRING" / "v12" / f"taxon_{args.taxid}"
    tax_dir.mkdir(parents=True, exist_ok=True)
    print(f"Resolving {len(genes)} input identifiers against STRING...", flush=True)
    original_mapping = string_mapping_query(genes, args.taxid, tax_dir)
    direct = {row["queryItem"]: row for row in original_mapping}
    explicit_alias: Path | None = None
    if getattr(args, "identifier_aliases", None):
        explicit_alias = Path(args.identifier_aliases)
        if not explicit_alias.is_file():
            raise ExpansionError(f"Identifier alias file not found: {explicit_alias}")
    discovered_alias_paths = discover_string_alias_paths(Path(args.gene_list), outdir, explicit_alias)
    alias_paths = [Path(args.gene_list), *discovered_alias_paths]
    if discovered_alias_paths:
        print("Identifier bridge source(s):", flush=True)
        for path in discovered_alias_paths:
            print(f"  - {path}", flush=True)
    else:
        print("No nearby GFF or gene-to-protein/UniProt alias table was found automatically.", flush=True)
    aliases = string_aliases(genes, alias_paths)
    candidates = list(dict.fromkeys(alias for gene in genes if gene not in direct for alias in aliases[gene]))
    alias_mapping = string_mapping_query(candidates, args.taxid, tax_dir) if candidates else []
    by_alias = {row["queryItem"]: row for row in alias_mapping}
    mapping_rows, unmapped = [], []
    for gene in genes:
        if gene in direct:
            mapping_rows.append({**direct[gene], "input_id": gene, "mapping_source": "submitted identifier"})
            continue
        matches = [by_alias[alias] for alias in aliases[gene] if alias in by_alias]
        ids = {row["stringId"] for row in matches}
        if len(ids) == 1:
            mapping_rows.append({**matches[0], "input_id": gene, "mapping_source": "verified annotation alias"})
        else:
            unmapped.append({"input_id": gene, "status": "ambiguous aliases" if ids else "unmapped", "aliases_tried": "; ".join(aliases[gene])})
    write_string_audit(outdir, mapping_rows, unmapped)
    if not mapping_rows:
        preview = ", ".join(genes[:8])
        raise ExpansionError(
            f"Organism validation passed: {species['official_name']} (taxonomy ID {args.taxid}). "
            f"STRING could not resolve any of the {len(genes)} submitted identifiers. Examples: {preview}.\n\n"
            "Locus tags may belong to a different strain/assembly or may be absent from STRING's aliases. "
            "The software searched the nearby DE/RNA-processing project tree but could not establish a verified protein mapping. "
            "Supply the matching gene_id-to-protein_id/UniProt table or GFF annotation in Identifier aliases, "
            "or paste supported protein identifiers in Manual input. Verify that the selected organism matches your reference. "
            "Changing the confidence score cannot fix an identifier mismatch. No species substitution was made.\n"
            f"The complete unmatched list is saved in {outdir / 'Intermediate files' / 'STRING unmapped identifiers.tsv'}."
        )
    mapped_queries = {row["input_id"] for row in mapping_rows}
    resolved_ids = list(dict.fromkeys(row["stringId"] for row in mapping_rows))
    print(f"Mapped {len(mapped_queries)}/{len(genes)} input genes to {len(resolved_ids)} STRING proteins; {len(unmapped)} unresolved. See identifier audit.", flush=True)
    mapping_raw = cache / "string_id_mapping.tsv"
    with mapping_raw.open("w", encoding="utf-8", newline="") as handle:
        fields = list(dict.fromkeys(key for row in mapping_rows for key in row))
        writer = csv.DictWriter(handle, delimiter="\t", fieldnames=fields)
        writer.writeheader(); writer.writerows(mapping_rows)
    common = {"identifiers": "\r".join(resolved_ids), "species": str(args.taxid), "caller_identity": "Bacterial_RNA_Analysis"}

    network_data = dict(common)
    network_data.update({"required_score": str(args.required_score), "network_type": args.network_type, "add_nodes": str(args.add_nodes)})
    network_key = hashlib.sha256(json.dumps({"resolved_ids": resolved_ids, "taxid": args.taxid, "required_score": args.required_score, "network_type": args.network_type, "add_nodes": args.add_nodes}, sort_keys=True).encode("utf-8")).hexdigest()[:20]
    network_library = tax_dir / f"network_{network_key}.tsv"
    if network_library.is_file() and network_library.stat().st_size > 0:
        network_text = network_library.read_text(encoding="utf-8", errors="replace")
        print(f"Using cached STRING network response from shared Database Library; live STRING retrieval skipped: {network_library}", flush=True)
    else:
        time.sleep(1)
        network_resp = request_with_retry("POST", f"{STRING_API}/tsv/network", data=network_data)
        network_text = network_resp.text
        network_library.write_text(network_text, encoding="utf-8")
        _library_manifest_entry(f"STRING network taxon {args.taxid} {network_key}", network_library, {"STRING_version": "12.0", "taxid": args.taxid, "network_type": args.network_type, "required_score": args.required_score})
        print(f"STRING network response saved in shared Database Library: {network_library}", flush=True)
    network_raw = cache / f"string_{args.network_type}_network.tsv"
    network_raw.write_text(network_text, encoding="utf-8")
    edges = parse_tsv_text(network_text)

    nodes: dict[str, dict[str, object]] = {}
    for row in mapping_rows:
        sid = row["stringId"]
        node = nodes.setdefault(sid, {"node": sid, "string_id": sid, "preferred_name": row.get("preferredName", ""), "input_ids": "", "taxon_id": args.taxid, "kind": "Network node"})
        node["input_ids"] = ";".join(filter(None, [str(node["input_ids"]), row["input_id"]]))
    normalized_edges: list[dict[str, object]] = []
    for row in edges:
        a = row.get("stringId_A") or row.get("preferredName_A") or row.get("protein1") or ""
        b = row.get("stringId_B") or row.get("preferredName_B") or row.get("protein2") or ""
        if not a or not b:
            continue
        nodes.setdefault(a, {"node": a, "string_id": row.get("stringId_A", ""), "preferred_name": row.get("preferredName_A", ""), "input_ids": "", "taxon_id": args.taxid, "kind": "Network node"})
        nodes.setdefault(b, {"node": b, "string_id": row.get("stringId_B", ""), "preferred_name": row.get("preferredName_B", ""), "input_ids": "", "taxon_id": args.taxid, "kind": "Network node"})
        normalized_edges.append({
            "source": a,
            "target": b,
            "edge_type": f"STRING_{args.network_type}",
            "combined_score": row.get("score", row.get("combined_score", "")),
            "nscore": row.get("nscore", ""),
            "fscore": row.get("fscore", ""),
            "pscore": row.get("pscore", ""),
            "ascore": row.get("ascore", ""),
            "escore": row.get("escore", ""),
            "dscore": row.get("dscore", ""),
            "tscore": row.get("tscore", ""),
            "STRING_version": "12.0",
        })

    expression_edges: list[dict[str, object]] = []
    if args.expression_edges:
        expression_edges = read_expression_edges(Path(args.expression_edges).resolve())
        input_to_string = {row["input_id"]: row["stringId"] for row in mapping_rows}
        for edge in expression_edges:
            for endpoint in ("source", "target"):
                original = str(edge[endpoint])
                edge["original_" + endpoint] = original
                edge[endpoint] = input_to_string.get(original, original)
    combined_edges = combined_network_evidence(normalized_edges, expression_edges)
    for node, attrs in nodes.items():
        preferred = str(attrs.get("preferred_name", "") or "").strip()
        submitted = str(attrs.get("input_ids", "") or "").strip()
        attrs["label"] = preferred or submitted or node
        attrs["genes"] = submitted or node

    # These tables drive the same offline, linked interactive network workspace
    # used by co-expression. GraphML remains available, but is no longer the
    # only practical way to inspect a STRING result.
    plot_edges = outdir / "network_edges.tsv"
    plot_nodes = outdir / "network_nodes.tsv"
    write_tsv(plot_edges, combined_edges)
    write_tsv(plot_nodes, list(nodes.values()))

    try:
        import networkx as nx
        graph = nx.Graph()
        for node, attrs in nodes.items():
            graph.add_node(node, **attrs)
        for edge in combined_edges:
            graph.add_edge(
                str(edge["source"]),
                str(edge["target"]),
                supporting_layers=int(edge.get("supporting_layers", 0)),
                expression_support=str(edge.get("expression_support", "no")),
                expression_weight=str(edge.get("expression_weight", "")),
                STRING_support=str(edge.get("STRING_support", "no")),
                STRING_combined_score=str(edge.get("STRING_combined_score", "")),
            )
        nx.write_graphml(graph, outdir / "Combined network evidence.graphml")
    except Exception:
        pass

    wb = Workbook(); wb.remove(wb.active)
    append_sheet(wb, "STRING edges", normalized_edges)
    append_sheet(wb, "Expression edges", expression_edges)
    append_sheet(wb, "Network edges", combined_edges)
    append_sheet(wb, "Network nodes", list(nodes.values()))
    append_sheet(wb, "Mapping audit", mapping_rows)
    append_sheet(wb, "Unmapped", unmapped)
    if not normalized_edges:
        print("Identifiers mapped successfully, but no edges met this network type and score. Review coverage or adjust the score; the mapped nodes and audit are retained.", flush=True)
    organism_name = species["official_name"] or species["STRING_name"]
    append_sheet(wb, "Network summary", [{"input_genes": len(genes), "mapped_queries": len(mapped_queries), "nodes": len(nodes), "edges": len(normalized_edges), "network_type": args.network_type, "required_score": args.required_score, "taxid": args.taxid, "organism": organism_name, "STRING_version": "12.0", "expression_edges": len(expression_edges), "combined_pairs": len(combined_edges)}])
    append_sheet(wb, "Provenance", [{"module": "STRING PPI", "software": "STRING REST API", "version": "12.0", "endpoint": STRING_API, "taxid": args.taxid, "organism": organism_name, "network_type": args.network_type, "required_score": args.required_score, "add_nodes": args.add_nodes, "expression_edge_input": args.expression_edges or "", "mapping_sha256": sha256_file(mapping_raw), "network_sha256": sha256_file(network_raw)}])
    style_workbook(wb)
    workbook = outdir / "PPI and network.xlsx"
    wb.save(workbook)
    interactive_report = outdir / "STRING PPI interactive.html"
    try:
        shared_python = Path(__file__).resolve().parents[2] / "Shared Downstream Components" / "Python"
        loaded: dict[str, object] = {}
        for module_name, filename in (
            ("bra_string_interactive_plots", "interactive_plots.py"),
            ("bra_string_visualization_studio", "visualization_studio.py"),
        ):
            module_path = shared_python / filename
            if not module_path.is_file():
                raise FileNotFoundError(module_path)
            spec = importlib.util.spec_from_file_location(module_name, module_path)
            if spec is None or spec.loader is None:
                raise RuntimeError(f"Could not load {module_path.name}")
            module = importlib.util.module_from_spec(spec)
            sys.modules[module_name] = module
            spec.loader.exec_module(module)
            loaded[filename] = module
        loaded["interactive_plots.py"].network_plots({
            "output_dir": str(outdir),
            "edge_file": str(plot_edges),
            "node_file": str(plot_nodes),
            "plots": ["gene_network"],
            "max_plot_edges": 1500,
            "show_node_labels": len(nodes) <= 80,
            "node_label_count": min(80, len(nodes)),
            "graph_layout": "spring",
            "method": "STRING",
        })
        if not getattr(args, "skip_standalone_report", False):
            loaded["visualization_studio.py"].build_static_report("network", outdir, destination=interactive_report)
            print(f"STRING_INTERACTIVE_REPORT\t{interactive_report}", flush=True)
    except Exception as exc:
        # Preserve the completed scientific workbook if an optional plotting
        # dependency fails, and put the exact report error in the live console.
        print(f"STRING INTERACTIVE WARNING: {type(exc).__name__}: {exc}", file=sys.stderr, flush=True)
    return {
        "status": "complete",
        "module": "string_network",
        "workbook": str(workbook),
        "mapping": str(outdir / "Intermediate files" / "STRING identifier mapping.tsv"),
        "interactive_report": str(interactive_report) if interactive_report.is_file() else "",
        "unmapped_audit": str(outdir / "Intermediate files" / "STRING unmapped identifiers.tsv"),
        "taxid": args.taxid,
        "organism": organism_name,
        "nodes": len(nodes),
        "STRING_edges": len(normalized_edges),
        "expression_edges": len(expression_edges),
        "combined_pairs": len(combined_edges),
        "unmapped": len(unmapped),
    }


def read_term2gene(path: Path, source: str) -> list[dict[str, str]]:
    """Read a simple TERM2GENE table or a compatible BioCyc/MetaCyc export.

    Recognized headers are intentionally broad because pathway exports often use
    UNIQUE-ID/PATHWAY and GENE/GENES rather than clusterProfiler's TERM/GENE.
    One row per pathway-gene pair is preferred; cells containing a comma,
    semicolon, or slash-delimited gene list are expanded automatically.
    """
    if not path.is_file():
        raise ExpansionError(f"Mapping file not found: {path}")
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        raw_rows = [["" if value is None else str(value) for value in row] for row in ws.iter_rows(values_only=True)]
    else:
        text = path.read_text(encoding="utf-8-sig", errors="replace")
        lines = [line for line in text.splitlines() if line.strip() and not line.lstrip().startswith("#")]
        if not lines:
            return []
        delimiter = "\t" if "\t" in lines[0] else ","
        raw_rows = [[str(value) for value in row] for row in csv.reader(lines, delimiter=delimiter)]
    raw_rows = [row for row in raw_rows if len(row) >= 2 and any(str(value).strip() for value in row)]
    if not raw_rows:
        return []
    headers = [_canonical_header(value) for value in raw_rows[0]]
    term_names = {"term", "termid", "pathway", "pathwayid", "uniqueid", "pathwayuniqueid", "frameid"}
    gene_names = {"gene", "genes", "geneid", "geneids", "locustag", "locustags", "protein", "proteinid", "accession"}
    label_names = {"termname", "pathwayname", "commonname", "name", "description"}
    term_idx = next((i for i, value in enumerate(headers) if value in term_names), None)
    gene_idx = next((i for i, value in enumerate(headers) if value in gene_names), None)
    label_idx = next((i for i, value in enumerate(headers) if value in label_names and i not in {term_idx, gene_idx}), None)
    has_header = term_idx is not None and gene_idx is not None
    if not has_header:
        term_idx, gene_idx, label_idx = 0, 1, 2 if len(raw_rows[0]) > 2 else None
    data_rows = raw_rows[1:] if has_header else raw_rows
    rows: list[dict[str, str]] = []
    for row in data_rows:
        if max(term_idx, gene_idx) >= len(row):
            continue
        term = str(row[term_idx]).strip()
        gene_cell = str(row[gene_idx]).strip()
        label = str(row[label_idx]).strip() if label_idx is not None and label_idx < len(row) else term
        if not term or not gene_cell:
            continue
        genes = [value.strip() for value in re.split(r"\s*[;,/]\s*", gene_cell) if value.strip()]
        for gene in genes:
            rows.append({"term_id": term, "gene_id": gene, "term_name": label or term, "source": source})
    if not rows:
        raise ExpansionError(
            f"No pathway-gene pairs were found in {path.name}. Use columns such as "
            "term_id/pathway_id, gene_id/locus_tag, and optional term_name/pathway_name."
        )
    return rows


def _parse_kegg_organism_catalog(text: str) -> list[dict[str, str]]:
    """Parse both KEGG list/organism and list/genome response formats."""
    rows: list[dict[str, str]] = []
    seen: set[tuple[str, str]] = set()
    for line in text.splitlines():
        fields = [field.strip() for field in line.split("\t")]
        genome = code = name = lineage = ""
        if len(fields) >= 3 and re.fullmatch(r"T\d{5,}", fields[0], flags=re.IGNORECASE):
            genome, code, name = fields[:3]
            lineage = fields[3] if len(fields) > 3 else ""
        elif len(fields) >= 2 and re.fullmatch(r"T\d{5,}", fields[0], flags=re.IGNORECASE):
            genome = fields[0]
            # list/genome represents code and name as either
            # "eco; Escherichia coli..." or "eco, Escherichia coli...".
            match = re.match(r"^([a-z][a-z0-9]{2,5})\s*[,;]\s*(.+)$", fields[1], flags=re.IGNORECASE)
            if match:
                code, name = match.group(1), match.group(2)
        if not genome or not code or not name:
            continue
        identity = (genome.casefold(), code.casefold())
        if identity in seen:
            continue
        seen.add(identity)
        rows.append({"genome": genome, "code": code, "name": name, "lineage": lineage})
    return rows


def _kegg_organism_catalog() -> list[dict[str, str]]:
    ensure_database_library()
    cache_dir = DATABASE_LIBRARY_ROOT / "KEGG"
    cache_dir.mkdir(parents=True, exist_ok=True)
    catalog_path = cache_dir / "organism_catalog.tsv"
    rows: list[dict[str, str]] = []
    if catalog_path.is_file() and catalog_path.stat().st_size > 0:
        text = catalog_path.read_text(encoding="utf-8", errors="replace")
        rows = _parse_kegg_organism_catalog(text)
        if rows:
            print(f"Using cached KEGG organism catalog from shared Database Library: {catalog_path}", flush=True)
            return rows

    # list/organism is documented by KEGG, but it can temporarily return HTTP
    # 400. list/genome is an equivalent documented source for T numbers,
    # organism codes and names, so prefer it and retain list/organism as fallback.
    failures: list[str] = []
    source_url = ""
    text = ""
    for endpoint in ("list/genome", "list/organism"):
        source_url = f"{KEGG_API}/{endpoint}"
        try:
            text = request_with_retry("GET", source_url).text
            rows = _parse_kegg_organism_catalog(text)
            if rows:
                break
            failures.append(f"{endpoint}: response contained no organism records")
        except ExpansionError as exc:
            failures.append(f"{endpoint}: {exc}")
            rows = []
    if not rows:
        raise ExpansionError("KEGG organism catalog could not be loaded. " + " | ".join(failures))
    normalized = "\n".join("\t".join((row["genome"], row["code"], row["name"], row["lineage"])) for row in rows) + "\n"
    catalog_path.write_text(normalized, encoding="utf-8")
    _library_manifest_entry("KEGG organism catalog", catalog_path, {"source": source_url, "records": len(rows)})
    return rows


def resolve_kegg_organism(query: str) -> tuple[str, str]:
    query = str(query or "").strip()
    if not query:
        raise ExpansionError("Enter a KEGG organism code, bacterial organism name, or NCBI taxonomy ID.")
    dropdown_match = re.match(r"^([A-Za-z][A-Za-z0-9]{2,5})\s*[·|]\s*", query)
    if dropdown_match:
        query = dropdown_match.group(1)
    if re.fullmatch(r"br\d+", query, flags=re.IGNORECASE):
        raise ExpansionError(
            f"{query} is a KEGG BRITE/table identifier, not an organism code. "
            "Choose a species from the Online KEGG organism list. For Amycolatopsis orientalis, use aori."
        )
    catalog = _kegg_organism_catalog()
    folded = query.casefold()
    direct = [row for row in catalog if folded in {row["code"].casefold(), row["genome"].casefold()}]
    if direct:
        return direct[0]["code"], direct[0]["name"]
    if query.isdigit() or folded.startswith("taxid:"):
        taxid = re.sub(r"(?i)^taxid:", "", query).strip()
        matches = []
        try:
            response = request_with_retry("GET", f"{KEGG_API}/list/genome/{urllib.parse.quote(taxid)}")
            requested_rows = _parse_kegg_organism_catalog(response.text)
            requested_genomes = {row["genome"].casefold() for row in requested_rows}
            matches = [row for row in catalog if row["genome"].casefold() in requested_genomes]
        except ExpansionError:
            matches = []
        if not matches:
            response = request_with_retry("GET", f"{KEGG_API}/link/genome/taxid:{urllib.parse.quote(taxid)}/species")
            genomes = {field.replace("gn:", "") for line in response.text.splitlines() for field in line.split("\t") if field.startswith("gn:")}
            matches = [row for row in catalog if row["genome"] in genomes]
    else:
        normalized = re.sub(r"[^a-z0-9]+", " ", folded).strip()
        matches = []
        for row in catalog:
            haystack = re.sub(r"[^a-z0-9]+", " ", f'{row["name"]} {row["lineage"]}'.casefold()).strip()
            if normalized and normalized in haystack:
                matches.append(row)
        exact = [row for row in matches if re.sub(r"[^a-z0-9]+", " ", row["name"].casefold()).strip() == normalized]
        if exact:
            matches = exact
    if not matches:
        raise ExpansionError(f"KEGG did not contain an organism matching '{query}'. Enter a KEGG code such as eco, a scientific name, or an NCBI taxonomy ID.")
    if len(matches) > 1:
        preview = ", ".join(f'{row["code"]} ({row["name"]})' for row in matches[:8])
        raise ExpansionError(f"'{query}' matches several KEGG organisms: {preview}. Enter the intended KEGG code for an unambiguous analysis.")
    return matches[0]["code"], matches[0]["name"]


def kegg_species(args: argparse.Namespace) -> dict[str, object]:
    code, name = resolve_kegg_organism(args.organism)
    genome = next((row["genome"] for row in _kegg_organism_catalog() if row["code"].casefold() == code.casefold()), "")
    print(f"Validated KEGG organism: {name} ({code}; {genome})", flush=True)
    return {"status": "complete", "module": "kegg_species", "query": args.organism, "code": code, "name": name, "genome": genome}


def kegg_term2gene(organism: str) -> tuple[list[dict[str, str]], str, str]:
    ensure_database_library()
    organism_code, organism_name = resolve_kegg_organism(organism)
    org = safe_name(organism_code, "organism")
    cache_dir = DATABASE_LIBRARY_ROOT / "KEGG" / org
    cache_dir.mkdir(parents=True, exist_ok=True)
    raw_path = cache_dir / "pathway_gene_links.tsv"
    names_path = cache_dir / "pathway_names.tsv"
    if raw_path.is_file() and raw_path.stat().st_size > 0:
        text = raw_path.read_text(encoding="utf-8", errors="replace")
        print(f"Using cached KEGG pathway mapping from shared Database Library; download skipped: {raw_path}", flush=True)
    else:
        response = request_with_retry("GET", f"{KEGG_API}/link/pathway/{urllib.parse.quote(organism_code)}")
        text = response.text
        raw_path.write_text(text, encoding="utf-8")
        _library_manifest_entry(f"KEGG {organism_code} pathway links", raw_path, {"organism": organism_code, "organism_name": organism_name, "source": KEGG_API})
        print(f"KEGG mapping downloaded once and saved for reuse in shared Database Library: {raw_path}", flush=True)
    if names_path.is_file() and names_path.stat().st_size > 0:
        names_text = names_path.read_text(encoding="utf-8", errors="replace")
    else:
        names_text = request_with_retry("GET", f"{KEGG_API}/list/pathway/{urllib.parse.quote(organism_code)}").text
        names_path.write_text(names_text, encoding="utf-8")
        _library_manifest_entry(f"KEGG {organism_code} pathway names", names_path, {"organism": organism_code, "source": KEGG_API})
    pathway_names = {fields[0].replace("path:", ""): fields[1] for line in names_text.splitlines() if len(fields := line.split("\t", 1)) == 2}
    rows = []
    for line in text.splitlines():
        if not line.strip() or "\t" not in line:
            continue
        gene, pathway = line.split("\t", 1)
        term_id = pathway.replace("path:", "")
        rows.append({"term_id": term_id, "gene_id": gene, "term_name": pathway_names.get(term_id, term_id), "source": f"KEGG ({organism_code})"})
    print(f"Resolved KEGG organism '{organism}' to {organism_code} ({organism_name}).", flush=True)
    return rows, organism_code, organism_name


def bh_adjust(pvals: list[float]) -> list[float]:
    n = len(pvals)
    order = sorted(range(n), key=lambda i: pvals[i])
    adjusted = [1.0] * n
    running = 1.0
    for rank_idx in range(n - 1, -1, -1):
        i = order[rank_idx]
        rank = rank_idx + 1
        value = min(running, pvals[i] * n / rank)
        running = value
        adjusted[i] = min(1.0, value)
    return adjusted


def ora_enrichment(selected: set[str], universe: set[str], mappings: list[dict[str, str]]) -> list[dict[str, object]]:
    term_genes: dict[str, set[str]] = defaultdict(set)
    names: dict[str, str] = {}
    sources: dict[str, str] = {}
    for row in mappings:
        gene = row["gene_id"]
        candidates = {gene, gene.split(":", 1)[-1]}
        matched = next((x for x in candidates if x in universe), None)
        if matched:
            term_genes[row["term_id"]].add(matched)
            names[row["term_id"]] = row.get("term_name", row["term_id"])
            sources[row["term_id"]] = row.get("source", "")
    M = len(universe)
    N = len(selected & universe)
    rows: list[dict[str, object]] = []
    pvals: list[float] = []
    for term, genes in term_genes.items():
        n = len(genes)
        overlap = sorted(selected & genes)
        k = len(overlap)
        if k == 0 or M == 0 or N == 0:
            continue
        p = float(hypergeom.sf(k - 1, M, n, N))
        pvals.append(p)
        rows.append({"source": sources.get(term, ""), "term_id": term, "term_name": names.get(term, term), "GeneRatio": f"{k}/{N}", "BgRatio": f"{n}/{M}", "Count": k, "pvalue": p, "genes": ",".join(overlap)})
    adj = bh_adjust(pvals)
    for row, value in zip(rows, adj):
        row["p_adjust_BH"] = value
    rows.sort(key=lambda r: (float(r["p_adjust_BH"]), float(r["pvalue"])))
    return rows


def pathway(args: argparse.Namespace) -> dict[str, object]:
    ensure_database_library()
    selected = set(read_gene_list(Path(args.gene_list), args.gene_column))
    universe = set(read_gene_list(Path(args.universe), args.universe_column)) if args.universe else set(selected)
    if not selected:
        raise ExpansionError("The selected gene list is empty.")
    if not universe:
        universe = set(selected)
    mappings: list[dict[str, str]] = []
    if args.term2gene:
        source = Path(args.term2gene).resolve(); _archive_mapping(source, "Custom TERM2GENE")
        mappings.extend(read_term2gene(source, "User mapping"))
    if args.biocyc_mapping:
        source = Path(args.biocyc_mapping).resolve(); _archive_mapping(source, "BioCyc")
        mappings.extend(read_term2gene(source, "BioCyc"))
    if args.metacyc_mapping:
        source = Path(args.metacyc_mapping).resolve(); _archive_mapping(source, "MetaCyc")
        mappings.extend(read_term2gene(source, "MetaCyc"))
    resolved_kegg_code = ""
    resolved_kegg_name = ""
    if args.kegg_organism:
        if not args.kegg_confirmed:
            raise ExpansionError("KEGG online mapping requires explicit --kegg-confirmed because KEGG REST/data have separate usage conditions.")
        kegg_rows, resolved_kegg_code, resolved_kegg_name = kegg_term2gene(args.kegg_organism)
        mappings.extend(kegg_rows)
        time.sleep(0.4)
    if not mappings:
        raise ExpansionError("No pathway mapping source was selected. Supply TERM2GENE/BioCyc/MetaCyc mapping, or enable KEGG with an organism code and explicit license confirmation.")

    results = ora_enrichment(selected, universe, mappings)
    outdir = Path(args.output_dir).resolve(); outdir.mkdir(parents=True, exist_ok=True)
    intermediate = outdir / "Intermediate files" / "Pathway annotation cache"; intermediate.mkdir(parents=True, exist_ok=True)
    enrichment_table = outdir / "enrichment_results.tsv"
    write_tsv(enrichment_table, results)
    write_tsv(intermediate / "TERM2GENE.tsv", mappings)
    if resolved_kegg_code:
        shared_kegg = DATABASE_LIBRARY_ROOT / "KEGG" / safe_name(resolved_kegg_code, "organism") / "pathway_gene_links.tsv"
        if shared_kegg.is_file():
            shutil.copy2(shared_kegg, intermediate / "KEGG pathway gene links.tsv")
    wb = Workbook(); wb.remove(wb.active)
    append_sheet(wb, "Enrichment results", results)
    append_sheet(wb, "Gene-pathway mapping", mappings)
    mapped = {r["gene_id"].split(":", 1)[-1] for r in mappings}
    append_sheet(wb, "Annotation coverage", [{"selected_genes": len(selected), "universe_genes": len(universe), "selected_mapped": len({g for g in selected if g in mapped}), "mapping_rows": len(mappings), "enriched_terms": len(results)}])
    sources = sorted({r.get("source", "") for r in mappings})
    append_sheet(wb, "Database summary", [{"source": s, "mapping_rows": sum(1 for r in mappings if r.get("source") == s)} for s in sources])
    append_sheet(wb, "Provenance", [{"module": "Functional and pathway enrichment", "engine": "Python hypergeometric ORA + BH", "software_version": VERSION, "KEGG_online": bool(args.kegg_organism), "KEGG_query": args.kegg_organism or "", "KEGG_organism_code": resolved_kegg_code, "KEGG_organism_name": resolved_kegg_name, "BioCyc_import": bool(args.biocyc_mapping), "MetaCyc_import": bool(args.metacyc_mapping), "TERM2GENE_import": bool(args.term2gene)}])
    style_workbook(wb)
    workbook = outdir / "Pathway enrichment.xlsx"; wb.save(workbook)
    interactive_report = outdir / "Pathway enrichment interactive.html"
    try:
        shared_python = Path(__file__).resolve().parents[2] / "Shared Downstream Components" / "Python"
        loaded: dict[str, object] = {}
        for module_name, filename in (
            ("bra_pathway_interactive_plots", "interactive_plots.py"),
            ("bra_pathway_visualization_studio", "visualization_studio.py"),
        ):
            module_path = shared_python / filename
            if not module_path.is_file():
                raise FileNotFoundError(module_path)
            spec = importlib.util.spec_from_file_location(module_name, module_path)
            if spec is None or spec.loader is None:
                raise RuntimeError(f"Could not load {module_path.name}")
            module = importlib.util.module_from_spec(spec)
            sys.modules[module_name] = module
            spec.loader.exec_module(module)
            loaded[filename] = module
        loaded["interactive_plots.py"].enrichment_plots({
            "output_dir": str(outdir),
            "enrichment_result_file": str(enrichment_table),
            "plots": ["bar", "dot", "enrichment_map", "gene_term"],
            "top_terms": 30,
            "network_term_limit": 20,
            "genes_per_term": 60,
        })
        if not getattr(args, "skip_standalone_report", False):
            loaded["visualization_studio.py"].build_static_report("enrichment", outdir, destination=interactive_report)
            print(f"PATHWAY_INTERACTIVE_REPORT\t{interactive_report}", flush=True)
    except Exception as exc:
        # Preserve the completed workbook and mappings even if an optional
        # plotting dependency is unavailable; the exact problem remains visible
        # in the module's live console.
        print(f"PATHWAY INTERACTIVE WARNING: {type(exc).__name__}: {exc}", file=sys.stderr, flush=True)
    return {
        "status": "complete", "module": "pathway", "workbook": str(workbook),
        "interactive_report": str(interactive_report) if interactive_report.is_file() else "",
        "terms": len(results), "mapping_rows": len(mappings),
    }


def parse_discovery_transcripts(path: Path) -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    with path.open("r", encoding="utf-8-sig", errors="replace") as fh:
        reader = csv.DictReader(fh, delimiter="\t")
        for row in reader:
            if not row:
                continue
            try:
                row["start"] = int(str(row.get("start", "0")))
                row["end"] = int(str(row.get("end", "0")))
            except ValueError:
                continue
            rows.append(dict(row))
    return rows


def build_tu_candidates(genes: list[Interval], transcripts: list[dict[str, object]]) -> list[dict[str, object]]:
    genes_by_key: dict[tuple[str, str], list[Interval]] = defaultdict(list)
    for g in genes:
        if g.strand in {"+", "-"}:
            genes_by_key[(g.seqid, g.strand)].append(g)
    for values in genes_by_key.values():
        values.sort(key=lambda g: g.start)
    tx_by_key: dict[tuple[str, str], list[Interval]] = defaultdict(list)
    for tx in transcripts:
        tx_by_key[(str(tx["seqid"]), str(tx["strand"]))].append(Interval(str(tx["seqid"]), int(tx["start"]), int(tx["end"]), str(tx["strand"]), str(tx.get("transcript_id", ""))))
    for values in tx_by_key.values(): values.sort(key=lambda x: x.start)

    output: list[dict[str, object]] = []
    idx = 0
    for key, gs in sorted(genes_by_key.items()):
        used: set[str] = set()
        for tx in tx_by_key.get(key, []):
            covered = [g for g in gs if overlap_bp(tx, g) >= max(1, int(g.length * 0.5))]
            if not covered:
                continue
            covered.sort(key=lambda g: g.start)
            idx += 1
            for g in covered: used.add(g.name)
            first, last = covered[0], covered[-1]
            if tx.strand == "+":
                leader = max(0, first.start - tx.start)
                trailer = max(0, tx.end - last.end)
            else:
                leader = max(0, tx.end - first.end)
                trailer = max(0, last.start - tx.start)
            output.append({
                "tu_id": f"BRA_TU_{idx:05d}",
                "seqid": tx.seqid,
                "start": tx.start,
                "end": tx.end,
                "strand": tx.strand,
                "genes": ",".join(g.name for g in covered),
                "gene_count": len(covered),
                "boundary_evidence": "coverage-inferred transcript boundary",
                "putative_5prime_leader_bp": leader,
                "putative_3prime_extension_bp": trailer,
                "model_prediction": "",
                "manual_assignment": "unreviewed",
                "locked": "no",
                "notes": "",
            })
        for g in gs:
            if g.name in used:
                continue
            idx += 1
            output.append({"tu_id": f"BRA_TU_{idx:05d}", "seqid": g.seqid, "start": g.start, "end": g.end, "strand": g.strand, "genes": g.name, "gene_count": 1, "boundary_evidence": "annotation only", "putative_5prime_leader_bp": 0, "putative_3prime_extension_bp": 0, "model_prediction": "", "manual_assignment": "unreviewed", "locked": "no", "notes": ""})
    return output


def run_prodigal(fasta: Path, outdir: Path) -> tuple[list[dict[str, object]], str]:
    exe = shutil.which("prodigal")
    if not exe:
        return [], "Prodigal not installed"
    gff = outdir / "prodigal_rbs_evidence.gff"
    proteins = outdir / "prodigal_predicted_proteins.faa"
    result = run_command([exe, "-i", str(fasta), "-a", str(proteins), "-f", "gff", "-o", str(gff), "-p", "single"])
    rows = []
    for iv in read_gff(gff):
        attrs = iv.attrs or {}
        rows.append({"seqid": iv.seqid, "start": iv.start, "end": iv.end, "strand": iv.strand, "predicted_id": iv.name, "start_type": attrs.get("start_type", ""), "rbs_motif": attrs.get("rbs_motif", ""), "rbs_spacer": attrs.get("rbs_spacer", ""), "rscore": attrs.get("rscore", ""), "sscore": attrs.get("sscore", "")})
    return rows, str(gff)


def tu_architecture(args: argparse.Namespace) -> dict[str, object]:
    root = Path(args.analysis_ready).resolve() if args.analysis_ready else None
    inferred = infer_analysis_ready_paths(root) if root else {}
    fasta = Path(args.fasta).resolve() if args.fasta else inferred.get("fasta")
    gff = Path(args.gff).resolve() if args.gff else inferred.get("gff")
    if not fasta or not Path(fasta).is_file() or not gff or not Path(gff).is_file():
        raise ExpansionError("TU Architecture requires reference FASTA and GFF3.")
    transcript_table = Path(args.transcripts).resolve() if args.transcripts else None
    if not transcript_table and root:
        possible = root / "Transcript discovery and architecture" / "predicted_transcripts.tsv"
        if possible.is_file(): transcript_table = possible
    if not transcript_table or not transcript_table.is_file():
        raise ExpansionError("TU Architecture requires predicted_transcripts.tsv from Transcript Discovery or an explicitly supplied transcript table.")
    genes = choose_genes(read_gff(Path(gff)))
    transcripts = parse_discovery_transcripts(transcript_table)
    tus = build_tu_candidates(genes, transcripts)
    outdir = Path(args.output_dir).resolve(); outdir.mkdir(parents=True, exist_ok=True)
    rbs_rows: list[dict[str, object]] = []
    notes: list[str] = []
    if args.prodigal:
        rbs_rows, note = run_prodigal(Path(fasta), outdir)
        notes.append(note)
    terminator_rows: list[dict[str, object]] = []
    if args.transterm:
        notes.append("TransTermHP is available as a specialized optional terminator engine. This build does not guess a caller command; supply verified output through --terminator-table when available.")
    if args.terminator_table:
        terminator_rows = import_generic_terminator_table(Path(args.terminator_table).resolve())
        notes.append(f"Imported {len(terminator_rows)} verified terminator evidence rows.")
    write_tsv(outdir / "manual_tu_review.tsv", tus)
    wb = Workbook(); wb.remove(wb.active)
    append_sheet(wb, "TU candidates", tus)
    append_sheet(wb, "RBS start evidence", rbs_rows)
    append_sheet(wb, "Terminator evidence", terminator_rows)
    append_sheet(wb, "Boundary evidence", [{"tu_id": r["tu_id"], "boundary_evidence": r["boundary_evidence"], "putative_5prime_leader_bp": r["putative_5prime_leader_bp"], "putative_3prime_extension_bp": r["putative_3prime_extension_bp"]} for r in tus])
    append_sheet(wb, "Manual review", tus)
    append_sheet(wb, "Provenance", [{"module": "Operon / TU Architecture", "version": VERSION, "input_transcripts": str(transcript_table), "reference_fasta": str(fasta), "annotation": str(gff), "prodigal": bool(args.prodigal), "transterm_requested": bool(args.transterm), "terminator_table": args.terminator_table or "", "notes": " | ".join(notes)}])
    style_workbook(wb)
    workbook = outdir / "Operon and transcription units.xlsx"; wb.save(workbook)
    return {"status": "complete", "module": "tu_architecture", "workbook": str(workbook), "tu_candidates": len(tus), "manual_review_tsv": str(outdir / "manual_tu_review.tsv"), "notes": notes}


def apply_tu(args: argparse.Namespace) -> dict[str, object]:
    table = Path(args.review_tsv).resolve()
    rows: list[dict[str, str]] = []
    with table.open("r", encoding="utf-8-sig", errors="replace") as fh:
        reader = csv.DictReader(fh, delimiter="\t")
        rows = [dict(r) for r in reader]
    accepted = [r for r in rows if str(r.get("manual_assignment", "")).lower() not in {"reject", "rejected", "exclude"}]
    outdir = Path(args.output_dir).resolve(); outdir.mkdir(parents=True, exist_ok=True)
    gff = outdir / "curated_transcription_units.gff3"
    bed = outdir / "curated_transcription_units.bed"
    with gff.open("w", encoding="utf-8", newline="\n") as gh, bed.open("w", encoding="utf-8", newline="\n") as bh:
        gh.write("##gff-version 3\n")
        for r in accepted:
            try: start, end = int(r["start"]), int(r["end"])
            except Exception: continue
            attrs = f"ID={safe_name(r.get('tu_id','TU'))};genes={r.get('genes','')};manual_assignment={r.get('manual_assignment','')};locked={r.get('locked','no')}"
            gh.write("\t".join([r.get("seqid", ""), "BacterialRNAAnalysis", "transcription_unit", str(start), str(end), ".", r.get("strand", "."), ".", attrs]) + "\n")
            bh.write("\t".join([r.get("seqid", ""), str(start-1), str(end), r.get("tu_id", "TU"), "0", r.get("strand", ".")]) + "\n")
    return {"status": "complete", "accepted_tus": len(accepted), "gff3": str(gff), "bed": str(bed)}


def check_tools(args: argparse.Namespace) -> dict[str, object]:
    ensure_database_library()
    tools = ["bedtools", "prodigal", "RNAfold", "cmscan", "transterm", "java"]
    rows = []
    for tool in tools:
        path = shutil.which(tool)
        rows.append({"tool": tool, "status": "ready" if path else "optional_missing", "path": path or ""})
    return {"status": "ok", "version": VERSION, "database_library": str(DATABASE_LIBRARY_ROOT), "tools": rows, "python_packages": {"requests": requests.__version__}}


def parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(description="Bacterial RNA Analysis scientific expansion backend")
    p.add_argument("--json", action="store_true", help="Print machine-readable JSON")
    sub = p.add_subparsers(dest="command", required=True)

    c = sub.add_parser("check")
    c.set_defaults(func=check_tools)

    d = sub.add_parser("discover")
    d.add_argument("--analysis-ready")
    d.add_argument("--fasta")
    d.add_argument("--gff")
    d.add_argument("--coverage-root")
    d.add_argument("--rockhopper-transcripts", help="Optional precomputed Rockhopper transcript table imported as independent evidence")
    d.add_argument("--output-dir", required=True)
    d.add_argument("--min-depth", type=float, default=1.0)
    d.add_argument("--max-gap", type=int, default=25)
    d.add_argument("--min-length", type=int, default=50)
    d.add_argument("--min-samples", type=int, default=2)
    d.add_argument("--min-antisense-overlap", type=int, default=30)
    d.add_argument("--srna-min", type=int, default=30)
    d.add_argument("--srna-max", type=int, default=500)
    d.add_argument("--rnafold", action="store_true")
    d.add_argument("--rfam", action="store_true")
    d.add_argument("--rfam-cm")
    d.add_argument("--rfam-clanin")
    d.add_argument("--threads", type=int, default=max(1, os.cpu_count() or 1))
    d.set_defaults(func=discover)

    s = sub.add_parser("string")
    s.add_argument("--gene-list", required=True)
    s.add_argument("--identifier-aliases", help="Verified gene_id to protein_id/UniProt aliases or matching GFF annotation")
    s.add_argument("--gene-column")
    s.add_argument("--expression-edges", help="Optional existing WGCNA/CEMiTool/GENIE3-style edge table with source/target columns")
    s.add_argument("--taxid", required=True, type=int)
    s.add_argument("--network-type", choices=("functional", "physical"), default="physical")
    s.add_argument("--required-score", type=int, default=700)
    s.add_argument("--add-nodes", type=int, default=0)
    s.add_argument("--output-dir", required=True)
    s.set_defaults(func=string_network)

    sv = sub.add_parser("string-species", help="Validate a taxonomy ID against the official STRING v12 organism catalog")
    sv.add_argument("--taxid", required=True, type=int)
    sv.set_defaults(func=string_species)

    kv = sub.add_parser("kegg-species", help="Resolve a KEGG organism code, name, T number, or NCBI taxonomy ID")
    kv.add_argument("--organism", required=True)
    kv.set_defaults(func=kegg_species)

    mw = sub.add_parser("manual-workbook", help="Materialize a saved guided Excel workbook as analysis input tables")
    mw.add_argument("--profile", required=True, choices=("de", "combined", "network", "ppi", "pathway", "transcript", "tu"))
    mw.add_argument("--action", choices=["extract", "editor-use", "editor-export", "editor-import"], default="extract")
    mw.add_argument("--workbook", required=True)
    mw.add_argument("--output-dir", required=True)
    mw.set_defaults(func=manual_workbook)

    e = sub.add_parser("pathway")
    e.add_argument("--gene-list", required=True)
    e.add_argument("--gene-column")
    e.add_argument("--universe")
    e.add_argument("--universe-column")
    e.add_argument("--term2gene")
    e.add_argument("--biocyc-mapping")
    e.add_argument("--metacyc-mapping")
    e.add_argument("--kegg-organism")
    e.add_argument("--kegg-confirmed", action="store_true")
    e.add_argument("--output-dir", required=True)
    e.set_defaults(func=pathway)

    t = sub.add_parser("tu")
    t.add_argument("--analysis-ready")
    t.add_argument("--fasta")
    t.add_argument("--gff")
    t.add_argument("--transcripts")
    t.add_argument("--prodigal", action="store_true")
    t.add_argument("--transterm", action="store_true")
    t.add_argument("--terminator-table", help="Optional verified/imported terminator calls in TSV/CSV form")
    t.add_argument("--output-dir", required=True)
    t.set_defaults(func=tu_architecture)

    a = sub.add_parser("apply-tu")
    a.add_argument("--review-tsv", required=True)
    a.add_argument("--output-dir", required=True)
    a.set_defaults(func=apply_tu)
    return p


def main() -> int:
    args = parser().parse_args()
    try:
        payload = args.func(args)
        print(json.dumps(payload, indent=2, ensure_ascii=False))
        return 0
    except ExpansionError as exc:
        payload = {"status": "error", "message": str(exc)}
        print(json.dumps(payload, indent=2, ensure_ascii=False), file=sys.stderr)
        return 2
    except Exception as exc:
        payload = {"status": "error", "message": f"{type(exc).__name__}: {exc}"}
        print(json.dumps(payload, indent=2, ensure_ascii=False), file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
