#!/usr/bin/env python3
"""Convert a UTF-8 OpDetect text report into a readable Word document."""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_BREAK, WD_LINE_SPACING
from docx.shared import Inches, Pt


def configure(document: Document, title: str) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(9.5)
    for style_name, size in (("Title", 22), ("Heading 1", 15), ("Heading 2", 12)):
        styles[style_name].font.name = "Aptos Display"
        styles[style_name].font.size = Pt(size)

    document.add_heading(title, level=0)
    subtitle = document.add_paragraph("Generated automatically by the OpDetect Windows RNA-seq pipeline")
    subtitle.runs[0].italic = True
    subtitle.runs[0].font.size = Pt(9)


def add_preformatted(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    run = paragraph.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(7.5)


def convert(source: Path, destination: Path, title: str) -> None:
    text = source.read_text(encoding="utf-8", errors="replace")
    document = Document()
    configure(document, title)

    in_console = False
    for raw_line in text.splitlines():
        line = raw_line.rstrip()
        if line.startswith("[") and line.endswith("]") and len(line) < 90:
            heading = line[1:-1].replace("_", " ").title()
            if heading == "Complete Console Output":
                document.add_page_break()
                in_console = True
            else:
                in_console = False
            document.add_heading(heading, level=1)
            continue

        if in_console:
            add_preformatted(document, line if line else " ")
            continue

        if not line:
            document.add_paragraph()
            continue

        if line.endswith(":") and "\t" not in line and len(line) < 100:
            paragraph = document.add_paragraph()
            run = paragraph.add_run(line)
            run.bold = True
            continue

        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(2)
        if "\t" in line or line.startswith("OPDETECT_PROGRESS"):
            run = paragraph.add_run(line)
            run.font.name = "Consolas"
            run.font.size = Pt(8)
        else:
            paragraph.add_run(line)

    footer = document.sections[0].footer.paragraphs[0]
    footer.alignment = 1
    footer_run = footer.add_run("OpDetect run record")
    footer_run.font.size = Pt(8)

    destination.parent.mkdir(parents=True, exist_ok=True)
    document.save(destination)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--title", default="OpDetect Run Report")
    args = parser.parse_args()
    convert(args.input, args.output, args.title)


if __name__ == "__main__":
    main()
