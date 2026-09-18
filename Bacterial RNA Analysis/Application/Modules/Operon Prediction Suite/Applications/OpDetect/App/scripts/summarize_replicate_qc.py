#!/usr/bin/env python3
"""Summarize FASTQ, alignment, coverage, and replicate agreement QC."""

from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

import numpy as np
import pandas as pd
from scipy.stats import spearmanr
from docx import Document
from docx.shared import Pt

MAX_CORRELATION_POSITIONS = 200_000


def nested(data: dict, *keys: str, default=np.nan):
    value = data
    for key in keys:
        if not isinstance(value, dict) or key not in value:
            return default
        value = value[key]
    return value


def parse_fastp(path: Path) -> dict[str, float]:
    if not path.is_file():
        return {}
    with path.open("r", encoding="utf-8") as handle:
        data = json.load(handle)
    return {
        "reads_before_fastp": nested(data, "summary", "before_filtering", "total_reads"),
        "reads_after_fastp": nested(data, "summary", "after_filtering", "total_reads"),
        "q30_before": nested(data, "summary", "before_filtering", "q30_rate"),
        "q30_after": nested(data, "summary", "after_filtering", "q30_rate"),
        "gc_after": nested(data, "summary", "after_filtering", "gc_content"),
        "duplication_rate": nested(data, "duplication", "rate"),
    }


def parse_hisat2(path: Path) -> dict[str, float]:
    if not path.is_file():
        return {}
    text = path.read_text(encoding="utf-8", errors="replace")
    match = re.search(r"([0-9]+(?:\.[0-9]+)?)% overall alignment rate", text)
    return {"hisat2_overall_alignment_pct": float(match.group(1)) if match else np.nan}


def parse_flagstat(path: Path) -> dict[str, float]:
    if not path.is_file():
        return {}
    text = path.read_text(encoding="utf-8", errors="replace")

    def first(pattern: str) -> float:
        match = re.search(pattern, text, flags=re.MULTILINE)
        return float(match.group(1)) if match else np.nan

    return {
        "total_alignments": first(r"^(\d+) \+ \d+ in total"),
        "mapped_alignments": first(r"^(\d+) \+ \d+ mapped"),
        "properly_paired_alignments": first(r"^(\d+) \+ \d+ properly paired"),
    }


def parse_idxstats(path: Path, replicate: str) -> list[dict[str, object]]:
    if not path.is_file():
        return []
    rows: list[dict[str, object]] = []
    table = pd.read_csv(
        path,
        sep="\t",
        header=None,
        names=["contig", "contig_length", "mapped_reads", "unmapped_reads"],
    )
    for row in table.itertuples(index=False):
        if str(row.contig) == "*":
            continue
        rows.append(
            {
                "replicate": replicate,
                "contig": str(row.contig),
                "contig_length": int(row.contig_length),
                "mapped_reads": int(row.mapped_reads),
                "unmapped_reads": int(row.unmapped_reads),
            }
        )
    return rows


def read_coverage(path: Path) -> tuple[dict[str, float], np.ndarray]:
    table = pd.read_csv(
        path,
        sep="\t",
        header=None,
        usecols=[2],
        names=["coverage"],
        dtype={"coverage": np.float32},
    )
    values = table["coverage"].to_numpy(dtype=np.float32)
    if values.size == 0:
        raise ValueError(f"Coverage file is empty: {path}")
    if values.size <= MAX_CORRELATION_POSITIONS:
        sampled = values
    else:
        indexes = np.linspace(
            0, values.size - 1, MAX_CORRELATION_POSITIONS, dtype=np.int64
        )
        sampled = values[indexes]
    summary = {
        "mean_depth": float(np.mean(values)),
        "median_depth": float(np.median(values)),
        "coverage_breadth_pct": float(np.mean(values > 0) * 100.0),
        "positions": int(values.size),
    }
    return summary, np.log1p(sampled.astype(np.float64))


def write_matrix(path: Path, names: list[str], matrix: np.ndarray) -> None:
    frame = pd.DataFrame(matrix, index=names, columns=names)
    frame.index.name = "replicate"
    frame.to_csv(path, sep="\t", float_format="%.5f")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--samples-tsv", required=True, type=Path)
    parser.add_argument("--log-dir", required=True, type=Path)
    parser.add_argument("--coverage-dir", required=True, type=Path)
    parser.add_argument("--coverage-prefix", default="base_cov")
    parser.add_argument("--output-dir", required=True, type=Path)
    args = parser.parse_args()

    samples = pd.read_csv(args.samples_tsv, sep="\t", dtype=str).fillna("")
    sample_column = samples.columns[0]
    names = [str(value) for value in samples[sample_column].tolist()]
    args.output_dir.mkdir(parents=True, exist_ok=True)

    rows: list[dict[str, object]] = []
    contig_rows: list[dict[str, object]] = []
    correlation_vectors: list[np.ndarray] = []
    for name in names:
        row: dict[str, object] = {"replicate": name}
        row.update(parse_fastp(args.log_dir / f"{name}.fastp.json"))
        row.update(parse_hisat2(args.log_dir / f"{name}.hisat2.log"))
        row.update(parse_flagstat(args.log_dir / f"{name}.flagstat.log"))
        contig_rows.extend(parse_idxstats(args.log_dir / f"{name}.idxstats.tsv", name))
        coverage_summary, vector = read_coverage(
            args.coverage_dir / f"{args.coverage_prefix}_{name}"
        )
        row.update(coverage_summary)
        rows.append(row)
        correlation_vectors.append(vector)

    qc = pd.DataFrame(rows)
    for column in ["q30_before", "q30_after", "gc_after", "duplication_rate"]:
        if column in qc.columns:
            qc[column] = pd.to_numeric(qc[column], errors="coerce") * 100.0
    qc.to_csv(args.output_dir / "sample-qc.tsv", sep="\t", index=False, float_format="%.5f")
    if contig_rows:
        pd.DataFrame(contig_rows).to_csv(
            args.output_dir / "contig-mapping.tsv", sep="\t", index=False
        )

    summary_lines = [
        "OpDetect biological replicate QC summary",
        "========================================",
        f"Real biological replicates: {len(names)}",
        "Replicates were processed separately; raw coverage was not averaged.",
        "",
    ]

    if len(names) >= 2:
        lengths = {len(vector) for vector in correlation_vectors}
        if len(lengths) != 1:
            raise SystemExit("Coverage vectors have inconsistent lengths")
        matrix = np.stack(correlation_vectors, axis=0)
        pearson = np.corrcoef(matrix)
        spearman = np.eye(len(names), dtype=float)
        for first in range(len(names)):
            for second in range(first + 1, len(names)):
                value = float(spearmanr(matrix[first], matrix[second]).statistic)
                spearman[first, second] = value
                spearman[second, first] = value
        write_matrix(args.output_dir / "replicate-pearson.tsv", names, pearson)
        write_matrix(args.output_dir / "replicate-spearman.tsv", names, spearman)

        upper = pearson[np.triu_indices(len(names), k=1)]
        median_pearson = float(np.nanmedian(upper))
        minimum_pearson = float(np.nanmin(upper))
        if median_pearson >= 0.90:
            agreement = "high"
        elif median_pearson >= 0.70:
            agreement = "moderate"
        else:
            agreement = "low"
        summary_lines.extend(
            [
                f"Median pairwise log1p coverage Pearson correlation: {median_pearson:.4f}",
                f"Minimum pairwise log1p coverage Pearson correlation: {minimum_pearson:.4f}",
                f"Descriptive replicate agreement: {agreement}",
                "The agreement category is a QC heuristic, not a universal biological cutoff.",
            ]
        )
    else:
        summary_lines.extend(
            [
                "Replicate correlation: unavailable because only one biological replicate was supplied.",
                "The model can run, but replicate-to-replicate robustness cannot be assessed.",
            ]
        )

    alignment_values = pd.to_numeric(
        qc["hisat2_overall_alignment_pct"]
        if "hisat2_overall_alignment_pct" in qc.columns
        else pd.Series([np.nan] * len(qc)),
        errors="coerce",
    )
    low_alignment = qc[alignment_values < 50]
    if not low_alignment.empty:
        summary_lines.append(
            "WARNING: One or more replicates had an overall HISAT2 alignment rate below 50%."
        )
    summary_lines.extend(
        [
            "",
            "Detailed files:",
            "  sample-qc.tsv",
            "  contig-mapping.tsv",
            "  replicate-pearson.tsv (when at least two replicates are present)",
            "  replicate-spearman.tsv (when at least two replicates are present)",
        ]
    )
    document = Document()
    document.add_heading("OpDetect Biological Replicate QC Summary", level=0)
    for line in summary_lines[2:]:
        if not line:
            document.add_paragraph()
        elif line.endswith(":"):
            paragraph = document.add_paragraph()
            run = paragraph.add_run(line)
            run.bold = True
        else:
            document.add_paragraph(line)
    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            run.font.name = "Aptos"
            if run.font.size is None:
                run.font.size = Pt(10)
    document.save(args.output_dir / "replicate-qc-summary.docx")
    print(f"Wrote replicate QC files to {args.output_dir}")
    print("\n".join(summary_lines))


if __name__ == "__main__":
    main()
