#!/usr/bin/env python3
"""Validate in-app input drafts and optionally import/export Excel workbooks.

The GUI holds pasted data in memory. Using a draft validates all required sheets
before materializing the TSV inputs consumed by existing analysis engines.
Pale examples are presentation-only placeholders inside the application and
never become analysis input. A real Excel workbook is written only when the
user chooses to export it.
"""

from __future__ import annotations

import argparse
import csv
import json
import math
import re
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


GREEN = "2A5534"
GREEN_SOFT = "E6F2E8"
INK = "1E2A22"
MUTED = "58655D"


PROFILES: dict[str, dict[str, Any]] = {
    "de": {
        "title": "Differential Expression manual input",
        "filename": "Manual Differential Expression input.xlsx",
        "sheets": [
            ("Raw counts", "raw_counts.tsv", True,
             ["gene_id", "Control_1", "Control_2", "Control_3", "Treatment_1", "Treatment_2", "Treatment_3"],
             "One gene per row. Sample columns must contain non-negative integer counts."),
            ("Sample metadata", "sample_metadata.tsv", True,
             ["sample_id", "condition", "batch"],
             "One row per sample. sample_id values must exactly match count-matrix columns."),
            ("Gene coordinates optional", "gene_coordinates.tsv", False,
             ["gene_id", "seqid", "start", "end", "strand"],
             "Optional genomic coordinates used for IGV and genome-region views."),
        ],
    },
    "combined": {
        "title": "Functional Enrichment and Co-expression manual input",
        "filename": "Manual Functional Enrichment and Co-expression input.xlsx",
        "sheets": [
            ("DE result", "differential_expression.tsv", True,
             ["gene_id", "baseMean", "log2FoldChange", "stat", "pvalue", "padj"],
             "Paste a complete differential-expression result. Keep a signed statistic for fgsea."),
            ("Normalized expression", "normalized_expression.tsv", True,
             ["gene_id", "Sample_1", "Sample_2", "Sample_3"],
             "One gene per row and one independent biological sample per remaining column."),
            ("Sample metadata", "sample_metadata.tsv", True,
             ["sample_id", "condition", "batch"],
             "sample_id values must exactly match normalized-expression columns."),
            ("Gene to term mapping optional", "gene_to_term_mapping.tsv", False,
             ["gene_id", "term_id", "term_name", "source"],
             "Optional when online gene-to-term mapping is used."),
            ("Universe optional", "gene_universe.tsv", False,
             ["gene_id"], "Optional custom tested-gene universe."),
            ("Regulators optional", "regulators.tsv", False,
             ["gene_id"], "Optional regulator list for GENIE3."),
        ],
    },
    "network": {
        "title": "Co-expression and Network manual input",
        "filename": "Manual Co-expression and Network input.xlsx",
        "sheets": [
            ("Normalized expression", "normalized_expression.tsv", True,
             ["gene_id", "Sample_1", "Sample_2", "Sample_3"],
             "One gene per row and one independent biological sample per remaining column."),
            ("Sample metadata", "sample_metadata.tsv", True,
             ["sample_id", "condition", "batch"],
             "sample_id values must exactly match normalized-expression columns."),
            ("Regulators optional", "regulators.tsv", False,
             ["gene_id"], "Optional regulator list for GENIE3."),
        ],
    },
    "ppi": {
        "title": "STRING PPI manual input",
        "filename": "Manual STRING PPI input.xlsx",
        "sheets": [
            ("Gene list", "gene_list.tsv", True, ["gene_id"],
             "Paste one STRING-recognizable gene, protein, or locus identifier per row."),
            ("Identifier aliases optional", "identifier_aliases.tsv", False,
             ["gene_id", "protein_id"],
             "Optional verified gene/locus-to-protein or UniProt identifiers from the same organism annotation."),
            ("Expression edges optional", "expression_edges.tsv", False,
             ["source", "target", "weight"],
             "Optional existing co-expression or regulatory edges."),
        ],
    },
    "pathway": {
        "title": "Pathway Database Analysis manual input",
        "filename": "Manual Pathway Database Analysis input.xlsx",
        "sheets": [
            ("Selected genes", "selected_genes.tsv", True, ["gene_id"],
             "Genes selected for pathway over-representation analysis."),
            ("Gene universe", "gene_universe.tsv", True, ["gene_id"],
             "All tested and mappable genes from the experiment."),
            ("TERM2GENE optional", "term2gene.tsv", False,
             ["term_id", "gene_id", "term_name"],
             "Optional local pathway-to-gene mapping; online KEGG can be used instead."),
            ("BioCyc optional", "biocyc_mapping.tsv", False,
             ["pathway_id", "gene_id", "pathway_name"], "Optional authorized BioCyc export."),
            ("MetaCyc optional", "metacyc_mapping.tsv", False,
             ["pathway_id", "gene_id", "pathway_name"], "Optional licensed MetaCyc export."),
        ],
    },
    "transcript": {
        "title": "Transcript Discovery manual input",
        "filename": "Manual Transcript Discovery input.xlsx",
        "path_rows": [
            ("RNA Processing result folder", "", "Preferred: completed RNA Processing Results folder"),
            ("Reference FASTA override", "", "Optional; contigs must match annotation and coverage"),
            ("Gene annotation override", "", "Optional GFF3/GFF/GTF"),
            ("Coverage folder override", "", "Optional folder containing strand-specific tracks"),
        ],
        "sheets": [
            ("Input paths", "input_paths.tsv", False,
             ["input_name", "file_or_folder_path", "notes"], "Enter Windows file or folder paths in the file_or_folder_path column."),
            ("Rockhopper optional", "rockhopper_transcripts.tsv", False,
             ["seqid", "start", "end", "strand", "transcript_id"],
             "Optional precomputed Rockhopper transcript evidence."),
        ],
    },
    "tu": {
        "title": "TU Architecture manual input",
        "filename": "Manual TU Architecture input.xlsx",
        "path_rows": [
            ("RNA Processing result folder", "", "Used to locate matching reference and annotation"),
            ("Reference FASTA override", "", "Optional"),
            ("Annotation override", "", "Optional GFF3/GFF/GTF"),
            ("Predicted transcripts file", "", "Optional alternative to pasting the worksheet below"),
        ],
        "sheets": [
            ("Input paths", "input_paths.tsv", False,
             ["input_name", "file_or_folder_path", "notes"], "Enter Windows file or folder paths in the file_or_folder_path column."),
            ("Predicted transcripts", "predicted_transcripts.tsv", False,
             ["transcript_id", "seqid", "start", "end", "strand", "class", "mean_depth"],
             "Paste Transcript Discovery output when not supplying a file path."),
            ("Terminator evidence optional", "terminator_evidence.tsv", False,
             ["seqid", "start", "end", "strand", "score", "source"],
             "Optional verified terminator calls."),
        ],
    },
}


def canonical(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(value).casefold())


def clean_cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).replace("\r", " ").replace("\n", " ").replace("\t", " ").strip()


def profile_config(profile: str) -> dict[str, Any]:
    key = profile.strip().casefold()
    if key not in PROFILES:
        raise ValueError(f"Unknown manual-workbook profile: {profile}")
    return PROFILES[key]


def style_sheet(ws) -> None:
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A2"
    for cell in ws[1]:
        cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=GREEN)
        cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 24
    for idx, cell in enumerate(ws[1], 1):
        ws.column_dimensions[get_column_letter(idx)].width = max(14, min(42, len(str(cell.value or "")) + 5))
    ws.auto_filter.ref = ws.dimensions


def create_workbook(profile: str, output: Path) -> Path:
    cfg = profile_config(profile)
    output = output.expanduser().resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    readme = wb.active
    readme.title = "README"
    readme.sheet_view.showGridLines = False
    readme["A1"] = cfg["title"]
    readme["A1"].font = Font(name="Calibri", size=18, bold=True, color=GREEN)
    readme.merge_cells("A1:D1")
    readme["A3"] = "How to use"
    readme["A3"].font = Font(bold=True, color=GREEN)
    instructions = [
        "1. Paste your data below the header row on each required worksheet. Keep gene IDs as text.",
        "2. Keep worksheet names and gene_id; rename/add/remove sample columns to match your experiment.",
        "3. Save and close the workbook in Excel.",
        "4. Open Manual Excel input, choose Import workbook, review the cells, and select Use data.",
        "5. Each import creates a fresh set of TSV input tables beside the workbook; optional blank sheets are not reused.",
    ]
    for row, text in enumerate(instructions, 4):
        readme.cell(row=row, column=1, value=text)
    readme.column_dimensions["A"].width = 118
    readme.column_dimensions["B"].width = 18

    path_rows = cfg.get("path_rows", [])
    for sheet_name, _filename, required, headers, description in cfg["sheets"]:
        ws = wb.create_sheet(sheet_name)
        ws.append(headers)
        if sheet_name == "Input paths":
            for row in path_rows:
                ws.append(list(row))
        for cell in ws[1]:
            cell.comment = Comment(description, "Bacterial RNA Analysis")
        style_sheet(ws)
        ws["A1"].comment = Comment(
            ("Required. " if required else "Optional. ") + description,
            "Bacterial RNA Analysis",
        )
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        if sheet_name == "Input paths":
            ws.column_dimensions["A"].width = 34
            ws.column_dimensions["B"].width = 72
            ws.column_dimensions["C"].width = 58

    wb.properties.title = cfg["title"]
    wb.properties.subject = "Editable analysis input template"
    wb.save(output)
    return output


def nonempty_rows(ws) -> list[list[str]]:
    rows: list[list[str]] = []
    for values in ws.iter_rows(values_only=True):
        row = [clean_cell(value) for value in values]
        while row and not row[-1]:
            row.pop()
        if row and any(row):
            rows.append(row)
    return rows


def find_sheet(workbook, requested: str):
    exact = {canonical(name): name for name in workbook.sheetnames}
    key = canonical(requested)
    if key in exact:
        return workbook[exact[key]]
    return None


def extract_workbook(profile: str, workbook_path: Path, output_dir: Path, *, editor_data: dict | None = None) -> dict[str, Any]:
    profile = profile.strip().casefold()
    cfg = profile_config(profile)
    workbook_path = workbook_path.expanduser().resolve()
    if editor_data is None and not workbook_path.is_file():
        raise FileNotFoundError(f"Manual input workbook not found: {workbook_path}")
    if editor_data is None and workbook_path.suffix.casefold() not in {".xlsx", ".xlsm"}:
        raise ValueError("Manual input must be an .xlsx or .xlsm workbook.")
    output_dir = output_dir.expanduser().resolve()
    if output_dir.exists() and any(output_dir.iterdir()):
        raise ValueError("Use a new or empty extraction folder; existing input tables will not be overwritten or reused.")
    wb = editor_workbook(profile, editor_data) if editor_data is not None else load_workbook(workbook_path, read_only=True, data_only=True)
    files: dict[str, str] = {}
    staged: dict[str, list[list[str]]] = {}
    missing: list[str] = []
    empty_required: list[str] = []
    try:
        for sheet_name, filename, required, headers, _description in cfg["sheets"]:
            ws = find_sheet(wb, sheet_name)
            if ws is None:
                if required:
                    missing.append(sheet_name)
                continue
            rows = nonempty_rows(ws)
            if not rows:
                if required:
                    empty_required.append(sheet_name)
                continue
            actual_headers = rows[0]
            matrix = sheet_name in {"Raw counts", "Normalized expression"}
            fixed_headers = ["gene_id"] if matrix else headers
            # Batch is optional; leaving its template column blank is also valid.
            fixed_headers = [value for value in fixed_headers if not (sheet_name == "Sample metadata" and value == "batch")]
            required_headers = [canonical(value) for value in fixed_headers]
            actual_canonical = [canonical(value) for value in actual_headers]
            if not all(actual_headers) or len(set(actual_canonical)) != len(actual_headers):
                raise ValueError(f"Worksheet '{sheet_name}' has empty or duplicate column headers.")
            absent = [fixed_headers[index] for index, key in enumerate(required_headers) if key not in actual_canonical]
            if absent:
                raise ValueError(f"Worksheet '{sheet_name}' is missing header(s): {', '.join(absent)}")
            data_rows = rows[1:]
            if matrix and (actual_canonical[0] != "geneid" or len(actual_headers) < 2):
                raise ValueError(f"Worksheet '{sheet_name}' needs gene_id in the first column and at least one sample column.")
            if any(len(row) > len(actual_headers) for row in data_rows):
                raise ValueError(f"Worksheet '{sheet_name}' has data beyond its header columns. Add the missing sample/field headers.")
            if sheet_name == "Input paths":
                path_index = actual_canonical.index(canonical("file_or_folder_path"))
                has_user_data = any(path_index < len(row) and bool(row[path_index].strip()) for row in data_rows)
            else:
                has_user_data = any(any(value for value in row) for row in data_rows)
            if required and not has_user_data:
                empty_required.append(sheet_name)
                continue
            if not has_user_data:
                continue
            # Preserve the documented field names even if users changed their case/spacing.
            header_names = {canonical(value): value for value in fixed_headers}
            normalized_headers = [header_names.get(canonical(value), value) for value in actual_headers]
            staged[filename] = [normalized_headers] + [row + [""] * (len(actual_headers) - len(row)) for row in data_rows]
    finally:
        wb.close()
    if missing:
        raise ValueError("Required worksheet(s) not found: " + ", ".join(missing))
    if empty_required:
        raise ValueError("Paste data below the header row in required worksheet(s): " + ", ".join(empty_required))
    if profile == "transcript" and "input_paths.tsv" not in staged:
        raise ValueError(
            "Transcript Discovery needs at least one file or folder path in the 'Input paths' worksheet. "
            "Normally enter the completed RNA Processing result folder."
        )
    if profile == "tu" and not ({"input_paths.tsv", "predicted_transcripts.tsv"} & set(staged)):
        raise ValueError(
            "TU Architecture needs at least one file/folder path or pasted rows in the "
            "'Predicted transcripts' worksheet."
        )
    # Validate complete tables before writing any outputs. Never partly assign a failed import.
    for matrix_file in ("raw_counts.tsv", "normalized_expression.tsv"):
        if matrix_file not in staged:
            continue
        table = staged[matrix_file]
        seen: set[str] = set()
        for row_number, row in enumerate(table[1:], 2):
            if not row[0] or row[0] in seen:
                raise ValueError(f"{matrix_file} row {row_number}: gene_id must be non-empty and unique.")
            seen.add(row[0])
            for column, value in zip(table[0][1:], row[1:]):
                try:
                    number = float(value)
                    valid = math.isfinite(number)
                    if matrix_file == "raw_counts.tsv":
                        valid = valid and number >= 0 and number.is_integer()
                except ValueError:
                    valid = False
                if not valid:
                    kind = "non-negative integer count" if matrix_file == "raw_counts.tsv" else "finite numeric value"
                    raise ValueError(f"{matrix_file} row {row_number}, sample '{column}': expected a {kind}.")
        metadata = staged.get("sample_metadata.tsv")
        if metadata:
            index = metadata[0].index("sample_id")
            samples = [row[index] for row in metadata[1:]]
            if not all(samples) or len(samples) != len(set(samples)):
                raise ValueError("Sample metadata needs a unique non-empty sample_id in every row.")
            if set(samples) != set(table[0][1:]):
                raise ValueError("Sample metadata sample_id values must exactly match the expression/count sample-column names.")
            condition_index = metadata[0].index("condition")
            if any(not row[condition_index] for row in metadata[1:]):
                raise ValueError("Sample metadata needs a condition for every sample.")
    output_dir.mkdir(parents=True, exist_ok=True)
    for filename, rows in staged.items():
        target = output_dir / filename
        with target.open("w", encoding="utf-8", newline="") as handle:
            writer = csv.writer(handle, delimiter="\t", lineterminator="\n")
            writer.writerows(rows)
        files[Path(filename).stem] = str(target)
    manifest = {
        "status": "complete",
        "profile": profile,
        "workbook": "" if editor_data is not None else str(workbook_path),
        "input_source": "in_app_editor" if editor_data is not None else "excel",
        "output_dir": str(output_dir),
        "files": files,
    }
    (output_dir / "manual_input_manifest.json").write_text(
        json.dumps(manifest, indent=2, ensure_ascii=False) + "\n", encoding="utf-8"
    )
    return manifest


def editor_schema(profile: str) -> dict[str, Any]:
    cfg = profile_config(profile)
    resource = Path(__file__).resolve().parents[1] / "Examples" / "Manual input workbooks" / f"{profile}.json"
    if resource.is_file():
        return json.loads(resource.read_text(encoding="utf-8"))
    return {"profile": profile, "title": cfg["title"], "sheets": [
        {"name": name, "headers": headers, "rows": list(cfg.get("path_rows", [])) if name == "Input paths" else [],
         "required": required, "description": description, "example_headers": headers, "example_rows": []}
        for name, _filename, required, headers, description in cfg["sheets"]]}


def editor_workbook(profile: str, data: dict):
    if data.get("profile") != profile:
        raise ValueError("This workbook belongs to a different analysis. Open its matching manual-input editor.")
    wb = Workbook()
    wb.remove(wb.active)
    allowed = {spec[0] for spec in profile_config(profile)["sheets"]}
    seen = set()
    for sheet in data.get("sheets", []):
        name = str(sheet.get("name", ""))
        if name not in allowed or name in seen:
            raise ValueError(f"Unexpected or duplicate input worksheet: {name}")
        seen.add(name)
        ws = wb.create_sheet(name)
        for values in [sheet.get("headers", [])] + sheet.get("rows", []):
            ws.append([clean_cell(v) for v in values])
            # Identifiers beginning with '=' are data, never executable spreadsheet formulas.
            for cell in ws[ws.max_row]:
                cell.data_type = "s"
    return wb


def export_editor(profile: str, data: dict, output: Path) -> dict:
    wb = editor_workbook(profile, data)
    readme = wb.create_sheet("README", 0)
    readme.append([profile_config(profile)["title"]])
    readme.append(["Input sheets are editable. Pale examples are displayed only inside the application and are not exported."])
    readme.append(["Paste below headers; sample columns can be renamed. Use data validates inputs without requiring an Excel export."])
    readme.column_dimensions["A"].width = 110
    for ws in wb.worksheets[1:]:
        style_sheet(ws)
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    return {"status": "complete", "workbook": str(output)}


def import_editor(profile: str, workbook: Path, output: Path) -> dict:
    result = editor_schema(profile)
    wb = load_workbook(workbook, read_only=True, data_only=True)
    try:
        for sheet in result["sheets"]:
            ws = find_sheet(wb, sheet["name"])
            rows = nonempty_rows(ws) if ws is not None else []
            if rows:
                sheet["headers"], sheet["rows"] = rows[0], rows[1:]
    finally:
        wb.close()
    output.write_text(json.dumps(result, ensure_ascii=False), encoding="utf-8")
    return {"status": "complete", "editor": str(output)}


def build_bundled(output_dir: Path) -> list[str]:
    output_dir.mkdir(parents=True, exist_ok=True)
    created = []
    for key, cfg in PROFILES.items():
        target = output_dir / cfg["filename"]
        create_workbook(key, target)
        created.append(str(target))
    return created


def parser() -> argparse.ArgumentParser:
    root = argparse.ArgumentParser(description="Editable Excel input workbook helper")
    commands = root.add_subparsers(dest="command", required=True)
    create = commands.add_parser("create")
    create.add_argument("--profile", required=True, choices=sorted(PROFILES))
    create.add_argument("--output", required=True)
    extract = commands.add_parser("extract")
    extract.add_argument("--profile", required=True, choices=sorted(PROFILES))
    extract.add_argument("--workbook", required=True)
    extract.add_argument("--output-dir", required=True)
    for action in ("editor-use", "editor-export", "editor-import"):
        command = commands.add_parser(action)
        command.add_argument("--profile", required=True, choices=sorted(PROFILES))
        command.add_argument("--input", required=True)
        command.add_argument("--output", required=True)
    bundled = commands.add_parser("build-bundled")
    bundled.add_argument("--output-dir", required=True)
    return root


def main() -> int:
    args = parser().parse_args()
    try:
        if args.command == "create":
            payload: Any = {"status": "complete", "workbook": str(create_workbook(args.profile, Path(args.output)))}
        elif args.command == "extract":
            payload = extract_workbook(args.profile, Path(args.workbook), Path(args.output_dir))
        elif args.command == "editor-use":
            data = json.loads(Path(args.input).read_text(encoding="utf-8-sig"))
            payload = extract_workbook(args.profile, Path(args.input), Path(args.output), editor_data=data)
        elif args.command == "editor-export":
            data = json.loads(Path(args.input).read_text(encoding="utf-8-sig"))
            payload = export_editor(args.profile, data, Path(args.output))
        elif args.command == "editor-import":
            payload = import_editor(args.profile, Path(args.input), Path(args.output))
        else:
            payload = {"status": "complete", "workbooks": build_bundled(Path(args.output_dir))}
        print(json.dumps(payload, indent=2, ensure_ascii=False))
        return 0
    except Exception as exc:
        print(json.dumps({"status": "error", "message": str(exc)}, ensure_ascii=False), file=__import__("sys").stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
