#!/usr/bin/env python3
"""Overlay transcriptome log2 fold changes on original KEGG PNG/KGML maps.

This is a pathway visualization, not an enrichment test or a flux estimate.
Only exact, organism-qualified gene IDs or explicit KO assignments are used.
Network access is confined to run(); importing the module does not fetch data.
"""
from __future__ import annotations

import argparse
import base64
import hashlib
import json
import math
import os
import re
import struct
import sys
import time
import urllib.error
import urllib.request
import xml.etree.ElementTree as ET
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import pandas as pd

API = "https://rest.kegg.jp"
VIEWER_NAME = "kegg_pathway_maps_interactive.html"
TABLES = [
    ("KEGG map summary", "kegg_map_summary.tsv"),
    ("KEGG mapped genes", "kegg_map_genes.tsv"),
    ("KEGG mapping audit", "kegg_map_audit.tsv"),
    ("KEGG map nodes", "kegg_map_nodes.tsv"),
]
MAX_PATHWAYS = 12
MAX_RESPONSE = 24 * 1024 * 1024


def truth(value: Any) -> bool:
    return str(value).strip().lower() in {"true", "1", "yes", "on"}


def local_path(value: Any) -> Path:
    text = str(value or "").strip()
    match = re.match(r"^([A-Za-z]):[\\/](.*)$", text)
    if match and os.name != "nt":
        text = f"/mnt/{match[1].lower()}/" + match[2].replace("\\", "/")
    return Path(text).expanduser()


def atomic_bytes(path: Path, data: bytes) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(path.name + ".tmp")
    temporary.write_bytes(data)
    temporary.replace(path)


def write_json(path: Path, data: Any) -> None:
    atomic_bytes(path, json.dumps(data, indent=2, ensure_ascii=False, allow_nan=False).encode("utf-8"))


def parse_pathways(value: Any, organism: str = "") -> list[str]:
    """Resolve generic map IDs to the selected organism, or to KO if none.

    Explicit organism pathway IDs must agree with the selected organism.
    This prevents a plausible-looking map assembled from the wrong species.
    """
    org = re.split(r"\s*[·|]\s*", organism.strip())[0].strip().lower()
    if org and not re.fullmatch(r"[a-z][a-z0-9]{1,5}", org):
        raise ValueError("Enter a KEGG organism code (for example eco), or leave it empty for KO maps.")
    if org in {"map", "ko", "ec", "rn", "br"}:
        raise ValueError("Choose a KEGG organism code, not a reference-map prefix.")
    raw = value if isinstance(value, list) else re.split(r"[\s,;]+", str(value or "").strip())
    result = []
    for item in raw:
        token = str(item).strip().lower().removeprefix("path:")
        if not token:
            continue
        match = re.fullmatch(r"([a-z][a-z0-9]*?)?(\d{5})", token)
        if not match:
            raise ValueError(f"Invalid KEGG pathway ID: {item}. Examples: 00010, eco00010, ko00010.")
        prefix, number = match.groups()
        if prefix in {"ec", "rn", "br"}:
            raise ValueError(f"{item} is not a gene or KO map. Use an organism pathway ID or ko{number}.")
        prefix = (org or "ko") if not prefix or prefix == "map" else prefix
        if org and prefix not in {org, "ko"}:
            raise ValueError(f"Pathway {item} belongs to {prefix}, but the selected organism is {org}.")
        canonical = prefix + number
        if canonical not in result:
            result.append(canonical)
    if not result or len(result) > MAX_PATHWAYS:
        raise ValueError(f"Enter between 1 and {MAX_PATHWAYS} KEGG pathway IDs, separated by commas or spaces.")
    return result


def column(frame: pd.DataFrame, names: tuple[str, ...]) -> str | None:
    lookup = {re.sub(r"[^a-z0-9]", "", str(c).lower()): str(c) for c in frame.columns}
    return next((lookup[n] for name in names if (n := re.sub(r"[^a-z0-9]", "", name.lower())) in lookup), None)


def read_table(path: Path, sheet: str = "", *, de: bool = False) -> pd.DataFrame:
    if not path.is_file():
        raise FileNotFoundError(f"Input table not found: {path}")
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        with pd.ExcelFile(path, engine="openpyxl") as book:
            if sheet:
                chosen = sheet
            elif de:
                lookup = {s.casefold(): s for s in book.sheet_names}
                chosen = next((lookup[s] for s in ("all contrast rows", "differential expression") if s in lookup), None)
                if chosen is None:
                    # Do not accidentally read a run-summary or significant-only sheet.
                    matches = []
                    for name in book.sheet_names:
                        if "significant" in name.lower():
                            continue
                        head = pd.read_excel(book, sheet_name=name, nrows=0)
                        if column(head, ("log2FoldChange", "logFC", "log2FC")) and column(head, ("gene_id", "gene", "locus_tag", "GeneID")):
                            matches.append(name)
                    if len(matches) != 1:
                        raise ValueError("Choose a DE worksheet with gene_id and log2FoldChange, or export the full contrast as TSV. Multiple candidate worksheets are ambiguous.")
                    chosen = matches[0]
            else:
                chosen = book.sheet_names[0]
            return pd.read_excel(book, sheet_name=chosen, dtype=str).fillna("")
    return pd.read_csv(path, sep="," if path.suffix.lower() == ".csv" else "\t", dtype=str, keep_default_na=False).fillna("")


def number(value: Any) -> float | None:
    try:
        result = float(value)
        return result if math.isfinite(result) else None
    except (ValueError, TypeError):
        return None


def tokens(value: Any) -> list[str]:
    return [s for s in re.split(r"[\s,;|/]+", str(value or "").strip()) if s and s.upper() not in {"NA", "NAN", "NONE", "-"}]


def input_records(config: dict[str, Any]) -> tuple[list[dict[str, Any]], list[str]]:
    frame = read_table(local_path(config.get("result_file")), str(config.get("kegg_map_result_sheet", "")), de=True)
    gene_col = column(frame, (str(config.get("result_gene_column", "gene_id")), "gene_id", "GeneID", "gene", "locus_tag"))
    lfc_col = column(frame, (str(config.get("lfc_column", "log2FoldChange")), "log2FoldChange", "logFC", "log2FC"))
    padj_col = column(frame, (str(config.get("padj_column", "padj")), "padj", "FDR", "adj.P.Val", "adjusted_p_value"))
    contrast_col = column(frame, ("contrast", "comparison", "contrast_name"))
    function_col = column(frame, ("product", "function", "description", "gene_function", "protein_name"))
    if gene_col is None or lfc_col is None:
        raise ValueError("The DE input must contain gene_id and log2FoldChange (or logFC/log2FC). A selected-gene list without fold changes cannot be mapped.")
    warnings = []
    if padj_col is None:
        warnings.append("No adjusted p-value column: expression changes are shown, but significance is unknown.")
    id_names = {"keggid", "kegggeneid", "ko", "koid", "keggko", "locustag", "uniprot", "uniprotid", "ncbigeneid", "proteinid"}
    id_cols = [c for c in frame.columns if re.sub(r"[^a-z0-9]", "", str(c).lower()) in id_names]
    bridge: dict[str, list[str]] = defaultdict(list)
    mapping = str(config.get("kegg_map_gene_mapping", "")).strip()
    if mapping:
        mapping_frame = read_table(local_path(mapping))
        key = column(mapping_frame, ("gene_id", "input_gene_id", "gene", "locus_tag"))
        targets = [c for c in mapping_frame.columns if re.sub(r"[^a-z0-9]", "", str(c).lower()) in {"keggid", "kegggeneid", "ko", "koid", "keggko", "locustag"} and c != key]
        if key is None or not targets:
            raise ValueError("The optional gene-ID mapping requires gene_id and kegg_id and/or ko_id columns. TERM2GENE is a different file format.")
        for row in mapping_frame.to_dict("records"):
            bridge[str(row[key]).strip()].extend(t for c in targets for t in tokens(row[c]))
    records, seen = [], set()
    for row in frame.to_dict("records"):
        gene = str(row[gene_col]).strip()
        if not gene:
            raise ValueError("The DE table contains an empty gene ID. Remove or correct that row before mapping.")
        contrast = str(row[contrast_col]).strip() if contrast_col else str(config.get("contrast_label", "DE contrast"))
        contrast = contrast or "DE contrast"
        if (gene, contrast) in seen:
            raise ValueError(f"Duplicate gene/contrast: {gene} / {contrast}. Supply one DE result per gene and contrast; values will not be averaged.")
        seen.add((gene, contrast))
        aliases = {gene, *bridge.get(gene, []), *(t for c in id_cols for t in tokens(row[c]))}
        # Prefix external IDs only when the column declares their namespace.
        for c in id_cols:
            kind = re.sub(r"[^a-z0-9]", "", str(c).lower())
            namespace = {"uniprot": "uniprot", "uniprotid": "uniprot", "ncbigeneid": "ncbi-geneid", "proteinid": "ncbi-proteinid"}.get(kind)
            if namespace:
                aliases.update(f"{namespace}:{t}" for t in tokens(row[c]) if ":" not in t)
                aliases.difference_update(t for t in tokens(row[c]) if t != gene and ":" not in t)
        padj = number(row[padj_col]) if padj_col else None
        if padj is not None and not 0 <= padj <= 1:
            raise ValueError(f"Adjusted p-value outside [0, 1] for {gene} / {contrast}.")
        records.append({"gene_id": gene, "contrast": contrast, "log2FoldChange": number(row[lfc_col]), "padj": padj,
                        "function": str(row[function_col]) if function_col else "", "ids": sorted(aliases)})
    if not records:
        raise ValueError("The DE input is empty.")
    return records, warnings


class KeggClient:
    """Small bounded REST client; successful, validated responses are cached."""
    def __init__(self, cache: Path, offline: bool = False, timeout: float = 25):
        self.cache, self.offline, self.timeout = cache, offline, timeout
        self.last_request = 0.0
        self.provenance: list[dict[str, Any]] = []

    def get(self, operation: str, filename: str, validate=None) -> bytes:
        path = self.cache / filename
        data = path.read_bytes() if path.is_file() else None
        cached = data is not None
        if data is not None and validate:
            try:
                validate(data)
            except Exception:
                data = None
                cached = False
        if data is None:
            if self.offline:
                raise FileNotFoundError(f"Cached KEGG file missing or invalid: {path}. Run once with cached-only mode cleared to download it.")
            last_error = None
            success = False
            for attempt in range(2):
                # KEGG requests are spaced below its documented 3 requests/sec.
                time.sleep(max(0, .4 - (time.monotonic() - self.last_request)))
                self.last_request = time.monotonic()
                try:
                    request = urllib.request.Request(f"{API}/{operation}", headers={"User-Agent": "BacterialRNAAnalysis/1.9.77 KEGGPathwayMaps"})
                    with urllib.request.urlopen(request, timeout=self.timeout) as response:
                        data = response.read(MAX_RESPONSE + 1)
                    if not data or len(data) > MAX_RESPONSE:
                        raise ValueError("KEGG returned an empty or oversized response.")
                    if validate:
                        validate(data)
                    atomic_bytes(path, data)
                    success = True
                    break
                except (OSError, ValueError, ET.ParseError) as exc:
                    last_error = exc
                    if isinstance(exc, urllib.error.HTTPError) and exc.code in {400, 404}:
                        break
                    if attempt == 0:
                        time.sleep(1)
            else:
                data = None
            if not success:
                raise RuntimeError(f"KEGG retrieval failed ({API}/{operation}): {last_error}") from last_error
        self.provenance.append({"url": f"{API}/{operation}", "cache_file": str(path), "cached": cached,
                                "sha256": hashlib.sha256(data).hexdigest(), "used_at_utc": datetime.now(timezone.utc).isoformat()})
        return data


def png_size(data: bytes) -> tuple[int, int]:
    if len(data) < 24 or not data.startswith(b"\x89PNG\r\n\x1a\n") or data[12:16] != b"IHDR":
        raise ValueError("KEGG image is not a PNG.")
    width, height = struct.unpack(">II", data[16:24])
    if not (0 < width <= 16000 and 0 < height <= 16000):
        raise ValueError("KEGG image dimensions are invalid.")
    from io import BytesIO
    from PIL import Image
    with Image.open(BytesIO(data)) as image:
        image.verify()
    return width, height


def parse_kgml(data: bytes, pathway: str) -> dict[str, Any]:
    if b"<!ENTITY" in data.upper():
        raise ValueError("KGML entity declarations are not supported.")
    root = ET.fromstring(data)
    if root.tag != "pathway" or root.get("name", "").removeprefix("path:") != pathway:
        raise ValueError(f"KGML pathway identity does not match {pathway}.")
    expected_org = re.sub(r"\d{5}$", "", pathway)
    if root.get("org") != expected_org:
        raise ValueError(f"KGML organism does not match pathway {pathway}.")
    entries = {e.get("id"): e for e in root.findall("entry")}
    def identifiers(entry, visiting):
        entry_id = entry.get("id")
        if entry_id in visiting:
            return set()
        if entry.get("type") == "group":
            return set().union(*(identifiers(entries[c.get("id")], visiting | {entry_id}) for c in entry.findall("component") if c.get("id") in entries))
        return set(entry.get("name", "").split()) if entry.get("type") in {"gene", "ortholog"} else set()
    nodes = []
    for entry_id, entry in entries.items():
        ids = sorted(identifiers(entry, set()))
        if not ids:
            continue
        for gi, graphic in enumerate(entry.findall("graphics")):
            kind = graphic.get("type", "rectangle")
            coords = [number(s) for s in graphic.get("coords", "").split(",")] if graphic.get("coords") else []
            if any(s is None for s in coords) or len(coords) % 2:
                coords = []
            x, y = number(graphic.get("x")), number(graphic.get("y"))
            width, height = number(graphic.get("width", 45)), number(graphic.get("height", 17))
            if kind == "line" and len(coords) >= 4:
                xs, ys = coords[0::2], coords[1::2]
                x, y, width, height = (min(xs)+max(xs))/2, (min(ys)+max(ys))/2, max(4,max(xs)-min(xs)), max(4,max(ys)-min(ys))
            if x is None or y is None or width is None or height is None:
                continue
            nodes.append({"id": f"{entry_id}.{gi}", "entry_id": entry_id, "type": entry.get("type"), "shape": kind,
                          "x": x, "y": y, "width": max(1,width), "height": max(1,height), "coords": coords,
                          "label": graphic.get("name", ""), "ids": ids})
    return {"id": pathway, "title": root.get("title", pathway), "organism": root.get("org", ""), "nodes": nodes}


def normalized_ids(record: dict[str, Any], organism: str, conversions: dict[str, set[str]]) -> set[str]:
    result = set()
    for value in record["ids"]:
        if re.fullmatch(r"(?:ko:)?K\d{5}", value):
            result.add("ko:" + value.removeprefix("ko:"))
        elif value.startswith(organism + ":") and organism != "ko":
            result.add(value)
        elif ":" not in value and organism != "ko":
            result.add(f"{organism}:{value}")
        result.update(conversions.get(value, set()))
    return result


def map_pathway(pathway: dict[str, Any], records: list[dict[str, Any]], conversions=None) -> tuple[list[dict], list[dict], list[dict]]:
    index: dict[str, set[int]] = defaultdict(set)
    for ri, record in enumerate(records):
        for value in normalized_ids(record, pathway["organism"], conversions or {}):
            index[value].add(ri)
    memberships: dict[int, list[str]] = defaultdict(list)
    node_rows, mapped, audit = [], [], []
    for node in pathway["nodes"]:
        hits = sorted(set().union(*(index.get(k, set()) for k in node["ids"])))
        node["records"] = hits
        for ri in hits:
            memberships[ri].append(node["id"])
        node_rows.append({"pathway_id": pathway["id"], "node_id": node["id"], "label": node["label"],
                          "kegg_ids": ";".join(node["ids"]), "shape": node["shape"], "x": node["x"], "y": node["y"],
                          "width": node["width"], "height": node["height"], "matched_genes": len({records[i]["gene_id"] for i in hits})})
    for ri, record in enumerate(records):
        common = {k: record[k] for k in ("gene_id", "contrast", "log2FoldChange", "padj", "function")}
        common.update(pathway_id=pathway["id"], node_ids=";".join(memberships.get(ri, [])))
        status = "matched" if ri in memberships else "not_on_pathway_or_unmapped"
        if ri in memberships and record["log2FoldChange"] is None:
            status = "matched_missing_fold_change"
        audit.append({**common, "status": status, "candidate_ids": ";".join(sorted(normalized_ids(record, pathway["organism"], conversions or {})))})
        if ri in memberships:
            mapped.append({**common, "status": status})
    return mapped, audit, node_rows


def write_tables(output: Path, summaries, mapped, audit, nodes) -> None:
    columns = [
        ["pathway_id", "contrast", "title", "status", "input_genes", "matched_genes", "coloured_genes", "unmatched_genes", "message"],
        ["pathway_id", "gene_id", "contrast", "log2FoldChange", "padj", "function", "node_ids", "status"],
        ["pathway_id", "gene_id", "contrast", "log2FoldChange", "padj", "function", "node_ids", "status", "candidate_ids"],
        ["pathway_id", "node_id", "label", "kegg_ids", "shape", "x", "y", "width", "height", "matched_genes"],
    ]
    for (_, filename), rows, names in zip(TABLES, (summaries,mapped,audit,nodes), columns):
        pd.DataFrame(rows, columns=names).to_csv(output / filename, sep="\t", index=False)


def update_workbook(output: Path) -> None:
    """Replace only this operation's sheets, retaining other analysis sheets."""
    from openpyxl import Workbook, load_workbook
    from finalize_results import write_frame, workbook_name
    target = output / workbook_name("combined", {})
    workbook = load_workbook(target) if target.is_file() else Workbook()
    if workbook.sheetnames == ["Sheet"] and workbook.active.max_row == 1:
        workbook.remove(workbook.active)
    for title, filename in TABLES:
        for old in list(workbook.sheetnames):
            if old == title or re.fullmatch(re.escape(title) + r" \d+", old):
                del workbook[old]
        frame = pd.read_csv(output / filename, sep="\t", keep_default_na=False)
        if frame.empty:
            workbook.create_sheet(title).append(list(frame.columns))
        else:
            write_frame(workbook, title, frame)
    temporary = target.with_name(target.stem + ".tmp.xlsx")
    workbook.save(temporary)
    workbook.close()
    temporary.replace(target)


def run(config: dict[str, Any], *, standalone: bool = False) -> dict[str, Any]:
    output = local_path(config["output_dir"])
    output.mkdir(parents=True, exist_ok=True)
    pathway_ids = parse_pathways(config.get("kegg_map_pathway_ids"), str(config.get("integrated_kegg_organism", "")))
    records, warnings = input_records(config)
    if any(not row["function"] for row in records):
        from visualization_studio import _member_annotation_lookup
        annotations = _member_annotation_lookup(output)
        for row in records:
            if not row["function"]:
                row["function"] = str(annotations.get(row["gene_id"], {}).get("function", ""))
    padj_cutoff, lfc_cutoff = number(config.get("padj_cutoff", .05)), number(config.get("lfc_cutoff", 1))
    if padj_cutoff is None or not 0 <= padj_cutoff <= 1 or lfc_cutoff is None or lfc_cutoff < 0:
        raise ValueError("Mapping requires a valid adjusted p-value cutoff [0,1] and a non-negative absolute log2 fold-change cutoff.")
    for row in records:
        row["significant"] = row["padj"] is not None and row["padj"] <= padj_cutoff and row["log2FoldChange"] is not None and abs(row["log2FoldChange"]) >= lfc_cutoff
    cache_root = local_path(config.get("kegg_map_cache_dir") or os.environ.get("BRA_DATABASE_LIBRARY") or (Path.home()/".local/share/prok-rnaseq/Database Library"))
    client = KeggClient(cache_root / "KEGG" / "Pathway maps", truth(config.get("kegg_map_offline", False)))
    contrasts = list(dict.fromkeys(r["contrast"] for r in records))
    pathways, summaries, mapped, audit, nodes = [], [], [], [], []
    conversion_cache: dict[str, dict[str, set[str]]] = {}
    for pid in pathway_ids:
        print(f"KEGG_MAP\t{pid}\tRetrieving native diagram and mapping measured genes", flush=True)
        try:
            kgml = client.get(f"get/{pid}/kgml", f"{pid}.kgml", lambda data: parse_kgml(data, pid))
            pathway = parse_kgml(kgml, pid)
            png = client.get(f"get/{pid}/image", f"{pid}.png", png_size)
            pathway["width"], pathway["height"] = png_size(png)
            pathway["image"] = "data:image/png;base64," + base64.b64encode(png).decode("ascii")
            pathway["url"] = f"https://www.kegg.jp/pathway/{pid}"
            org = pathway["organism"]
            if org not in conversion_cache:
                conversions: dict[str, set[str]] = defaultdict(set)
                if org != "ko":
                    namespaces = {v.split(":", 1)[0] for r in records for v in r["ids"] if v.startswith(("uniprot:","ncbi-geneid:","ncbi-proteinid:"))}
                    for namespace in sorted(namespaces):
                        try:
                            lines = client.get(f"conv/{org}/{namespace}", f"{org}_{namespace}.tsv").decode("utf-8").splitlines()
                            for line in lines:
                                pair = line.split("\t")
                                if len(pair) == 2:
                                    a,b=pair
                                    if a.startswith(org+":"): a,b=b,a
                                    if b.startswith(org+":"): conversions[a].add(b)
                        except Exception as exc:
                            warnings.append(f"{pid}: external ID conversion unavailable; exact gene IDs still used. {exc}")
                conversion_cache[org] = conversions
            m, a, n = map_pathway(pathway, records, conversion_cache[org])
            mapped.extend(m); audit.extend(a); nodes.extend(n)
            pathways.append(pathway)
            for contrast in contrasts:
                total = sum(r["contrast"] == contrast for r in records)
                hits = [r for r in m if r["contrast"] == contrast]
                summaries.append({"pathway_id":pid,"contrast":contrast,"title":pathway["title"],"status":"mapped" if hits else "no_matches",
                                  "input_genes":total,"matched_genes":len(hits),"coloured_genes":sum(r["log2FoldChange"] is not None for r in hits),
                                  "unmatched_genes":total-len(hits),"message":"" if hits else "No exact gene/KO matches. Check organism and the optional gene-ID mapping. Unmatched genes are not assumed unchanged."})
        except Exception as exc:
            message = str(exc)
            warnings.append(f"{pid}: {message}")
            summaries.append({"pathway_id":pid,"contrast":"","title":"","status":"unavailable","input_genes":len(records),"matched_genes":0,"coloured_genes":0,"unmatched_genes":"","message":message})
            audit.extend({"pathway_id":pid,**{k:r[k] for k in ("gene_id","contrast","log2FoldChange","padj","function")},"node_ids":"","status":"pathway_unavailable","candidate_ids":";".join(r["ids"])} for r in records)
    payload = {"schema_version":1,"pathways":pathways,"records":records,"contrasts":contrasts,"summary":summaries,"warnings":warnings,
               "padj_cutoff":padj_cutoff,"lfc_cutoff":lfc_cutoff,"generated_at":datetime.now(timezone.utc).isoformat(),"provenance":client.provenance,
               "source_table":str(local_path(config.get("result_file"))),"source_sha256":hashlib.sha256(local_path(config.get("result_file")).read_bytes()).hexdigest()}
    write_tables(output, summaries, mapped, audit, nodes)
    technical = output / "Intermediate files" / "KEGG pathway maps"
    write_json(technical / "mapping data.json", payload)
    template = Path(__file__).with_name("kegg_pathway_viewer.html").read_text(encoding="utf-8")
    # JSON is data, never executable HTML; neutralize script-closing text in annotations.
    encoded = json.dumps(payload, ensure_ascii=False, allow_nan=False).replace("<", "\\u003c").replace("\u2028","\\u2028").replace("\u2029","\\u2029")
    atomic_bytes(output / "Figures" / VIEWER_NAME, template.replace("__KEGG_MAP_DATA__", encoded).encode("utf-8"))
    if standalone:
        update_workbook(output)
        from visualization_studio import build_static_report
        build_static_report("combined", output)
    print(f"KEGG_MAP_RESULT\t{len(pathways)}/{len(pathway_ids)} diagrams available; {len(mapped)} mapped gene/contrast rows", flush=True)
    for warning in warnings:
        print(f"WARNING\t{warning}", flush=True)
    return payload


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("config", type=Path)
    parser.add_argument("--standalone", action="store_true", help="Update the combined workbook and interactive report without running GO or network analyses")
    args = parser.parse_args()
    try:
        payload = run(json.loads(args.config.read_text(encoding="utf-8-sig")), standalone=args.standalone)
        return 0 if payload["pathways"] else 2
    except Exception as exc:
        print(f"ERROR: KEGG pathway mapping: {exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
