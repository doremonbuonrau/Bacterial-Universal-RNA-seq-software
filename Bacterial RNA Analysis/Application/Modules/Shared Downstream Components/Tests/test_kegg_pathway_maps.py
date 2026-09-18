"""Regression tests with synthetic KGML/PNG fixtures; no network required."""
from __future__ import annotations

import io
import json
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "Python"))
import kegg_pathway_maps as maps


def fixture(directory: Path) -> dict:
    from PIL import Image, ImageDraw
    root = directory / "cache" / "KEGG" / "Pathway maps"
    root.mkdir(parents=True, exist_ok=True)
    xml = '''<?xml version="1.0"?>
<pathway name="path:eco00010" org="eco" title="Synthetic mapping test — not a biological pathway">
 <entry id="1" name="eco:b0001 eco:b0002" type="gene"><graphics name="A + B" x="240" y="150" width="140" height="38" type="rectangle"/></entry>
 <entry id="2" name="eco:b0003" type="gene"><graphics name="C" x="500" y="150" width="120" height="38" type="rectangle"/></entry>
 <entry id="3" name="cpd:C00001" type="compound"><graphics name="Compound" x="380" y="265" width="30" height="30" type="circle"/></entry>
 <entry id="4" name="group" type="group"><component id="1"/><component id="2"/><graphics name="Complex" x="680" y="310" width="150" height="38" type="roundrectangle"/></entry>
 <entry id="5" name="eco:b0004" type="gene"><graphics name="D" x="200" y="310" width="120" height="38" type="circle"/></entry>
 <entry id="6" name="eco:b0003" type="gene"><graphics name="C line" type="line" coords="430,310,485,355,540,310"/></entry>
</pathway>'''
    (root/"eco00010.kgml").write_text(xml)
    image = Image.new("RGB", (900,470), "white")
    draw = ImageDraw.Draw(image)
    draw.text((30,25), "SYNTHETIC TEST DIAGRAM - NOT A BIOLOGICAL PATHWAY", fill="#273d31")
    draw.line([(310,150),(440,150)], fill="#81958a", width=2)
    draw.line([(500,169),(500,225),(380,225),(380,250)], fill="#81958a", width=2)
    for cx,cy,w,h,label in [(240,150,140,38,"A + B"),(500,150,120,38,"C"),(680,310,150,38,"Complex")]:
        draw.rectangle((cx-w/2,cy-h/2,cx+w/2,cy+h/2), outline="#415b4c", width=1)
        draw.text((cx-15,cy-5),label,fill="#222")
    draw.ellipse((140,291,260,329),outline="#415b4c")
    draw.text((195,305),"D",fill="#222")
    draw.ellipse((365,250,395,280),outline="#555")
    draw.line([(430,310),(485,355),(540,310)], fill="#415b4c",width=2)
    image.save(root/"eco00010.png")
    (root/"ko00010.kgml").write_text(xml.replace('path:eco00010','path:ko00010').replace('org="eco"','org="ko"').replace('type="gene"','type="ortholog"').replace('eco:b0001','ko:K00001').replace('eco:b0002','ko:K00002').replace('eco:b0003','ko:K00003').replace('eco:b0004','ko:K00004'))
    (root/"ko00010.png").write_bytes((root/"eco00010.png").read_bytes())
    result = directory/"de.tsv"
    result.write_text('gene_id\tcontrast\tlog2FoldChange\tpadj\tproduct\n'
                     'b0001\tTreatment vs Control\t2\t0.001\tSynthetic enzyme A\n'
                     'b0002\tTreatment vs Control\t-3\t0.003\tSynthetic enzyme B\n'
                     'local_C\tTreatment vs Control\t0\t0.8\tSynthetic enzyme C\n'
                     'b0004\tTreatment vs Control\tNA\tNA\tMissing value example\n'
                     'unmatched_gene\tTreatment vs Control\t1\t0.01\tNot on test map\n'
                     'b0001\tRecovery vs Control\t-1.2\t0.01\tSynthetic enzyme A\n'
                     'b0002\tRecovery vs Control\t1.5\t0.2\tSynthetic enzyme B\n'
                     'local_C\tRecovery vs Control\t0.5\t0.02\tSynthetic enzyme C\n')
    bridge=directory/'bridge.tsv'
    bridge.write_text('gene_id\tkegg_id\tko_id\nlocal_C\teco:b0003\tK00003\nb0001\teco:b0001\tK00001\nb0002\teco:b0002\tK00002\nb0004\teco:b0004\tK00004\n')
    return {'result_file':str(result),'output_dir':str(directory/'results'),'integrated_kegg_organism':'eco',
            'kegg_map_pathway_ids':'00010,ko00010,00020','kegg_map_gene_mapping':str(bridge),
            'kegg_map_cache_dir':str(directory/'cache'),'kegg_map_offline':True,'padj_cutoff':.05,'lfc_cutoff':1}


class MappingTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory()
        self.root=Path(self.tmp.name)
        self.config=fixture(self.root)

    def tearDown(self):
        self.tmp.cleanup()

    def test_pathway_and_organism_validation(self):
        self.assertEqual(maps.parse_pathways('00010,map00020,eco00010,ko00030','eco · E. coli'),['eco00010','eco00020','ko00030'])
        self.assertEqual(maps.parse_pathways('map00010',''),['ko00010'])
        for pid,org in [('hsa00010','eco'),('../eco00010','eco'),('ec00010','eco'),('','eco')]:
            with self.assertRaises(ValueError): maps.parse_pathways(pid,org)

    def test_mapping_keeps_contrasts_opposite_signs_missing_and_unmatched(self):
        with patch('urllib.request.urlopen',side_effect=AssertionError('Offline must never access the network')):
            payload=maps.run(self.config)
        self.assertEqual(len(payload['pathways']),2)
        eco=payload['pathways'][0]
        first=next(n for n in eco['nodes'] if n['id']=='1.0')
        values=[payload['records'][i]['log2FoldChange'] for i in first['records']]
        self.assertEqual(values,[2,-3,-1.2,1.5])
        self.assertEqual((first['x'],first['y'],first['width']),(240,150,140))
        self.assertFalse(any(n['entry_id']=='3' for n in eco['nodes']))
        group=next(n for n in eco['nodes'] if n['id']=='4.0')
        self.assertEqual(len(group['records']),6)
        audit=maps.pd.read_csv(Path(self.config['output_dir'])/'kegg_map_audit.tsv',sep='\t')
        self.assertTrue((audit.loc[audit.gene_id=='unmatched_gene','status'].isin(['not_on_pathway_or_unmapped','pathway_unavailable'])).all())
        missing=audit[(audit.gene_id=='b0004')&(audit.pathway_id=='eco00010')].iloc[0]
        self.assertEqual(missing.status,'matched_missing_fold_change')
        self.assertTrue(maps.pd.isna(missing.log2FoldChange))
        zero=audit[(audit.gene_id=='local_C')&(audit.pathway_id=='eco00010')].iloc[0]
        self.assertEqual(zero.log2FoldChange,0)
        self.assertEqual(len(payload['warnings']),1)  # Missing cache for eco00020.
        self.assertEqual(len(payload['provenance']),4)

    def test_duplicate_values_are_rejected(self):
        source=Path(self.config['result_file'])
        with source.open('a') as f:f.write('b0001\tTreatment vs Control\t4\t0.01\tduplicate\n')
        with self.assertRaisesRegex(ValueError,'Duplicate gene/contrast'):maps.input_records(self.config)

    def test_namespace_and_ko_matching(self):
        row={'ids':['hsa:b0001','ko:K00001','K00002','uniprot:P12345']}
        self.assertEqual(maps.normalized_ids(row,'eco',{'uniprot:P12345':{'eco:b0003'}}),{'ko:K00001','ko:K00002','eco:b0003'})
        self.assertEqual(maps.normalized_ids({'ids':['b0001']},'ko',{}),set())

    def test_workbook_selection_and_preserving_existing_results(self):
        from openpyxl import Workbook,load_workbook
        wb=Workbook();wb.active.title='Run summary';wb.active.append(['Not the data'])
        ws=wb.create_sheet('Differential expression');ws.append(['gene_id','log2FoldChange','padj']);ws.append(['b0001',2,.01])
        long=wb.create_sheet('All contrast rows');long.append(['gene_id','contrast','log2FoldChange','padj']);long.append(['b0001','A vs B',1,.01]);long.append(['b0001','C vs B',-1,.02])
        source=self.root/'DE.xlsx';wb.save(source)
        cfg={**self.config,'result_file':str(source)}
        rows,_=maps.input_records(cfg)
        self.assertEqual([r['contrast'] for r in rows],['A vs B','C vs B'])
        maps.run(cfg)
        out=Path(cfg['output_dir']);original=Workbook();original.active.title='Enrichment results';original.active.append(['ID','geneID']);original.active.append(['GO:example','b0001'])
        name=out/'Functional enrichment and co-expression results.xlsx';original.save(name)
        maps.update_workbook(out);maps.update_workbook(out)
        with name.open('rb') as f:
            saved=load_workbook(f)
            self.assertEqual(saved['Enrichment results']['A2'].value,'GO:example')
            self.assertEqual(len([s for s in saved.sheetnames if s.startswith('KEGG ')]),4)

    def test_cache_reuse_and_corrupt_response_fail_closed(self):
        cache=self.root/'fetch';client=maps.KeggClient(cache)
        data=(self.root/'cache/KEGG/Pathway maps/eco00010.png').read_bytes()
        with patch('urllib.request.urlopen',return_value=io.BytesIO(data)) as fetch:
            self.assertEqual(client.get('get/eco00010/image','map.png',maps.png_size),data)
            self.assertEqual(client.get('get/eco00010/image','map.png',maps.png_size),data)
            self.assertEqual(fetch.call_count,1)
        (cache/'map.png').write_bytes(b'bad image')
        with patch('urllib.request.urlopen',side_effect=OSError('network unavailable')),patch('time.sleep'):
            with self.assertRaises(RuntimeError):client.get('get/eco00010/image','map.png',maps.png_size)

    def test_html_treats_gene_annotations_as_data(self):
        source=Path(self.config['result_file'])
        source.write_text(source.read_text().replace('Synthetic enzyme A','</script><script>window.BAD=true</script>'))
        maps.run(self.config)
        html=(Path(self.config['output_dir'])/'Figures'/maps.VIEWER_NAME).read_text()
        self.assertNotIn('</script><script>window.BAD',html)
        self.assertIn('\\u003c/script>',html)

    def test_combined_coordinator_and_finalization_keep_the_map(self):
        import integrated_external_analysis as integrated
        import finalize_results as final
        from openpyxl import load_workbook
        out=Path(self.config['output_dir']);out.mkdir()
        (out/'selected_genes_used.tsv').write_text('gene_id\nb0001\n')
        (out/'gene_universe_used.tsv').write_text('gene_id\nb0001\nb0002\n')
        (out/'enrichment_results.tsv').write_text('ID\tgeneID\tp.adjust\nGO:test\tb0001/b0002\t0.01\n')
        (out/'network_edges.tsv').write_text('source\ttarget\tweight\nb0001\tb0002\t0.9\n')
        cfg={**self.config,'kegg_map_enabled':True,'kegg_map_pathway_ids':'00010',
             'integrated_pathway_enabled':False,'integrated_string_enabled':False}
        config_path=out/'Intermediate files/combined configuration.json';config_path.parent.mkdir()
        config_path.write_text(json.dumps(cfg))
        with patch.object(integrated,'load_backend',side_effect=AssertionError('Unrequested external analyses must not load')):
            result=integrated.run(config_path)
        self.assertEqual(result['status'],'complete')
        target=final.create_workbook('combined',cfg,out)
        final.organize_files('combined',out,config_path)
        with target.open('rb') as f:
            wb=load_workbook(f)
            self.assertIn('KEGG mapped genes',wb.sheetnames)
            self.assertIn('Enrichment results',wb.sheetnames)
            self.assertEqual(wb['KEGG mapped genes'].max_row,8)
        self.assertTrue((out/'Figures'/maps.VIEWER_NAME).is_file())
        self.assertTrue((out/'Intermediate files/KEGG pathway maps/mapping data.json').is_file())


if __name__=='__main__':unittest.main(verbosity=2)
